"""Compare record matching on a public CSV and explicitly artificial derivatives.

The matcher receives only generated Excel files, never planted correspondences.
The original public CSV is preserved byte-for-byte. This evaluates matching,
not provenance classification, audit conclusions, or Excel recalculation.
"""
import argparse
from collections import Counter, defaultdict
import csv
from decimal import Decimal
import hashlib
import html
import json
from pathlib import Path
import random
import sys
from time import perf_counter

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
LAB = HERE.parents[1]
sys.path.insert(0, str(LAB))
from workbench.content import compare_content, typed_value
from workbench.cli import save_workbook_analysis
from workbench.content_report import write_content_report


def header_baselines(left, right, count):
    """Independent small-table baselines, with the same 80% row threshold.

    Exact whole-record comparison includes all columns. Fuzzy row comparison
    uses header correspondence and columns with >=3 distinct nontrivial values,
    like the value-based matcher. Both abstain on tied row assignments.
    """
    lhs = {c.value: c.column for c in left[1]}
    rhs = {c.value: c.column for c in right[1]}
    mapping = [(col, rhs[name]) for name, col in lhs.items() if name in rhs]
    if len(mapping) != left.max_column or len(mapping) != right.max_column:
        return {'exact': [], 'fuzzy': [], 'scope': 'Not applicable: complete header correspondence unavailable'}
    def key(sheet, row, side):
        return tuple(typed_value(sheet.cell(row, pair[side])) for pair in mapping)
    a, b = defaultdict(list), defaultdict(list)
    for row in range(2, count+2):
        a[key(left, row, 0)].append(row)
        b[key(right, row, 1)].append(row)
    exact = [(rows[0], b[token][0]) for token, rows in a.items()
             if len(rows) == 1 and len(b.get(token, [])) == 1]
    nontrivial = lambda t: t is not None and t[0] != 'bool' and t not in (('number', 0), ('number', 1), ('number', -1))
    eligible = []
    for lc, rc in mapping:
        values = [{typed_value(sheet.cell(r, col)) for r in range(1, count+2)}
                  for sheet, col in ((left, lc), (right, rc))]
        if all(sum(nontrivial(t) for t in tokens) >= 3 for tokens in values):
            eligible.append((lc, rc))
    scores = {}
    for ar in range(2, count+2):
        for br in range(2, count+2):
            equal = [typed_value(left.cell(ar, ac)) for ac, bc in eligible
                     if typed_value(left.cell(ar, ac)) is not None
                     and typed_value(left.cell(ar, ac)) == typed_value(right.cell(br, bc))]
            if eligible and len(equal)/len(eligible) >= .8 and len({t for t in equal if nontrivial(t)}) >= 2:
                scores[ar, br] = len(equal)
    best_a, best_b = {}, {}
    for row in range(2, count+2):
        for side, best in ((0, best_a), (1, best_b)):
            options = [(score, pair[1-side]) for pair, score in scores.items() if pair[side] == row]
            if options:
                high = max(score for score, _ in options)
                winners = [other for score, other in options if score == high]
                if len(winners) == 1:
                    best[row] = winners[0]
    fuzzy = [(ar, br) for ar, br in best_a.items() if best_b.get(br) == ar]
    return {'exact': exact, 'fuzzy': fuzzy, 'scope': 'Headers supply column correspondence; exact or mutually unique 80% row agreement'}


def generate(folder, headers, rows, case, changed_row):
    count, width = len(rows), len(headers)
    rng = random.Random(20260914)
    row_order = list(range(count)); rng.shuffle(row_order)
    col_order = list(range(width)); rng.shuffle(col_order)
    source, derivative = Workbook(), Workbook()
    a, b = source.active, derivative.active
    a.title, b.title = 'Published data', 'Transformed data'
    a.append(headers)
    for row in rows: a.append(row)
    renamed = case == 'renamed_headers_changed_amount'
    b.append([f'Field {i+1}' if renamed else headers[c] for i, c in enumerate(col_order)])
    marginal = case == 'columns_shuffled_independently'
    column_orders = []
    for c in range(width):
        order = list(range(count)); random.Random(20260914+c).shuffle(order)
        column_orders.append(order)
    truth = []
    for dest, origin in enumerate(row_order):
        values = []
        for col in col_order:
            src_row = column_orders[col][dest] if marginal else origin
            value = rows[src_row][col]
            if not marginal and origin == changed_row and col == 7:
                value = float(Decimal(str(value)) + Decimal('123.45'))
            values.append(value)
        b.append(values)
        if not marginal: truth.append((origin+2, dest+2))
    # A visible example total is our addition, not a formula in the publisher's CSV.
    # It is separated from the input table and its value is not evaluated here.
    for book, title, col in ((source, a.title, 8), (derivative, b.title, col_order.index(7)+1)):
        total = book.create_sheet('Review total')
        total['A1'] = 'Added demonstration SUM; not in published CSV'
        letter = get_column_letter(col)
        total['A2'] = f"=SUM('{title}'!{letter}2:{letter}{count+1})"
    folder.mkdir(parents=True)
    paths = [folder/'published.xlsx', folder/'derived.xlsx']
    for path, book in zip(paths, (source, derivative)): book.save(path)
    return paths, [source, derivative], truth, col_order


def metrics(pairs, truth):
    observed, expected = set(map(tuple, pairs)), set(map(tuple, truth))
    return {'returned': len(observed), 'correct_planted_pairs': len(observed & expected),
            'outside_planted_pairs': len(observed-expected), 'missed': len(expected-observed)}


def run(output):
    output = Path(output)
    if output.exists() and any(output.iterdir()): raise ValueError('Output folder must be new or empty')
    output.mkdir(parents=True, exist_ok=True)
    source = json.loads((HERE/'SOURCES.json').read_text())
    csv_path = HERE/'HMT_spending_over_25000_Sep_25.csv'
    assert hashlib.sha256(csv_path.read_bytes()).hexdigest() == source['sha256']
    with csv_path.open(encoding='utf-8-sig', newline='') as handle: raw = list(csv.reader(handle))
    headers = [h.strip() for h in raw[0]]
    rows = [row[:7] + [float(Decimal(row[7].strip().removeprefix('£').replace(',', '')))] + row[8:] for row in raw[1:]]
    vouchers = Counter(row[6] for row in rows)
    # Predetermined identifiable example, plus a separate repeated-voucher case
    # to expose the ambiguity rather than silently choosing a favorable record.
    single = next(i for i, row in enumerate(rows) if vouchers[row[6]] == 1)
    cases = [('reordered', None), ('changed_unique_voucher', single),
             ('changed_repeated_voucher', 0), ('renamed_headers_changed_amount', single),
             ('columns_shuffled_independently', None)]
    observations = []
    for case, changed in cases:
        paths, books, truth, col_order = generate(output/case/'inputs', headers, rows, case, changed)
        baseline_start = perf_counter()
        baseline = header_baselines(books[0].active, books[1].active, len(rows))
        baseline_ms = (perf_counter()-baseline_start)*1000
        start = perf_counter(); result = compare_content(paths, books); elapsed = (perf_counter()-start)*1000
        blocks = [b for b in result['blocks'] if b.get('kind') == 'row_alignment'
                  and b['left']['sheet'] == 'Published data' and b['right']['sheet'] == 'Transformed data']
        pairs = [(r['left'], r['right']) for b in blocks for r in b['row_mapping'] if r['left'] > 1 and r['right'] > 1]
        changed_pair = next((pair for pair in truth if pair[0] == changed+2), None) if changed is not None else None
        item = {'case': case, 'records': len(rows), 'changed_source_row': changed+2 if changed is not None else None,
                'row_alignment': metrics(pairs, truth), 'header_exact': metrics(baseline['exact'], truth),
                'header_fuzzy': metrics(baseline['fuzzy'], truth), 'baseline_scope': baseline['scope'],
                'changed_pair_returned': changed_pair in pairs if changed_pair else None,
                'ambiguous_rows_omitted': result['record_alignment']['ambiguous_rows_omitted'],
                'total_displayed_regions': len(result['blocks']),
                'informational_single_run_ms': {'complete_content_search': elapsed, 'both_header_baselines': baseline_ms}}
        observations.append(item)
        (output/case/'content.json').write_text(json.dumps(result, indent=2)+'\n')
        (output/case/'evaluation-only-truth.json').write_text(json.dumps({'row_pairs': truth, 'destination_column_order_zero_based': col_order}, indent=2)+'\n')
        print(json.dumps(item), flush=True)
    for case in ('changed_unique_voucher', 'changed_repeated_voucher', 'renamed_headers_changed_amount'):
        report = save_workbook_analysis(list((output/case/'inputs').glob('*.xlsx')), output/case/'report')
        content_path = report.with_name('content.json')
        content = json.loads(content_path.read_text())
        content['example_context'] = 'Örnek veri: HM Treasury Eylül 2025 harcama CSV’sinden üretilen iki Excel dosyası. Sıra/başlık değişiklikleri, +123,45 tutar değişikliği ve SUM formülleri bu deney için eklendi. Bunlar yayımlanmış verideki hata veya denetim bulgusu değildir.'
        content_path.write_text(json.dumps(content, indent=2, ensure_ascii=False)+'\n')
        write_content_report(content, report)
    receipt = {'scope': 'Public spending values, artificially transformed Excel files. Five engineering cases; no field accuracy or provenance truth claim.',
               'source': source, 'records': len(rows), 'preprocessing': 'UTF-8 BOM removed by CSV decoder; headers stripped; Amount parsed as numeric after removing leading £ and thousands commas; all other fields preserved as text. Demonstration SUM added to each workbook.',
               'transformation': 'Fixed seed 20260914. Row/column permutations. Changed value +123.45. Changed unique voucher is first unique voucher; repeated-voucher case changes first record. Renamed headers are Field N. Last case independently shuffles every column: identical marginal bags do not imply matching records.',
               'evaluation': 'Planted maps are separate files and are never passed to the matcher. Counts exclude header rows. Outside planted pairs in the marginal-shuffle case are coincidental row candidates, not claimed false provenance. Timing is informational, not a performance comparison.',
               'source_hashes': {str(p.relative_to(LAB)): hashlib.sha256(p.read_bytes()).hexdigest() for p in (LAB/'workbench/content.py', LAB/'workbench/record_matches.py', LAB/'workbench/cli.py', LAB/'workbench/content_report.py', Path(__file__))},
               'cases': observations}
    (output/'receipt.json').write_text(json.dumps(receipt, indent=2)+'\n')
    def cell(m): return f"{m['correct_planted_pairs']} doğru · {m['outside_planted_pairs']} diğer · {m['missed']} bulunamadı"
    table = ''.join('<tr><th>'+html.escape(r['case'])+'</th>'+''.join('<td>'+('Uygulanamadı: başlıklar eşlenemedi' if k!='row_alignment' and r['baseline_scope'].startswith('Not applicable') else cell(r[k]))+'</td>' for k in ('row_alignment', 'header_exact', 'header_fuzzy'))+'</tr>' for r in observations)
    (output/'index.html').write_text('''<!doctype html><html lang="tr"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Lab · Gerçek veri üzerinde kayıt eşleştirme</title><style>body{font:17px/1.6 system-ui;margin:40px auto;max-width:1100px;padding:0 20px;color:#182c36}table{border-collapse:collapse;width:100%;font-size:14px}th,td{padding:14px;border:1px solid #ccd8dd;text-align:left}a{color:#08696c}code{background:#eef3f5}aside{background:#fff3d2;padding:20px}</style><h1>84 harcama kaydı: sıra değişince ne bulunuyor?</h1><p>Veri: <a href="'''+html.escape(source['publication'], quote=True)+'''">HM Treasury, Eylül 2025</a>. Kaynak CSV korundu. Aşağıdaki Excel dosyaları, sıralama değişiklikleri, tutar farkı ve SUM formülleri bu deney için üretildi; yayımlanmış dosyalardaki hata veya bir denetim bulgusu değildir.</p><p><a href="changed_unique_voucher/report/index.html">Değişen tutarı ve bağlı formülleri incele →</a></p><table><thead><tr><th>Deney</th><th>Değerlerden sütun + satır eşleme</th><th>Başlıklarla tam kayıt eşitliği</th><th>Başlıklarla %80 satır eşleme</th></tr></thead><tbody>'''+table+'''</tbody></table><p>“Doğru” yalnız deneyde kaydedilen satır haritasına göre doğrudur. Son deneyde her sütun ayrı karıştırıldı: tüm sütun değer kümeleri aynı kaldı, fakat kayıt haritası kurulmadı. Buradaki “diğer” adaylar tesadüfi benzerlikleri gösterir.</p><aside>Bu eşleştirici bağımsız kanıt sayısını, kopyanın yönünü veya özgün kaynağı belirlemez. Aynı faturaya ait benzer satırlarda belirsiz eşlemeyi atlayabilir. Başlıklar korunduğunda basit satır karşılaştırması güçlü bir alternatiftir; bu beş örnek bilimsel yenilik ya da saha doğruluğu ölçümü değildir.</aside><p><a href="receipt.json">Ölçüm ve kaynak kaydı</a> · <a href="../SOURCES.json">Kaynak ve lisans</a></p></html>''')
    return output/'index.html'


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True)
    print(run(parser.parse_args().output))
