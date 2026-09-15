"""Conservative static XLSX lineage. Never computes or certifies spreadsheet values."""
from collections import defaultdict
import hashlib
from pathlib import Path
import re
import zipfile
from .contracts import DependencyProblem

DIRECT = re.compile(r'^\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?$')
DYNAMIC = {'INDIRECT', 'OFFSET', 'LAMBDA', 'LET'}
KNOWN_FUNCTIONS = {'SUM', 'AVERAGE', 'MIN', 'MAX', 'COUNT', 'COUNTA', 'COUNTBLANK', 'COUNTIF', 'COUNTIFS',
                   'SUMIF', 'SUMIFS', 'AVERAGEIF', 'AVERAGEIFS', 'IF', 'IFS', 'IFERROR', 'IFNA',
                   'AND', 'OR', 'NOT', 'XOR', 'ABS', 'ROUND', 'ROUNDUP', 'ROUNDDOWN', 'INT', 'MOD',
                   'POWER', 'SQRT', 'PRODUCT', 'SUMPRODUCT', 'SUBTOTAL', 'AGGREGATE',
                   'INDEX', 'MATCH', 'VLOOKUP', 'HLOOKUP', 'XLOOKUP', 'XMATCH', 'ROWS', 'COLUMNS',
                   'CONCAT', 'CONCATENATE', 'TEXTJOIN', 'TEXT', 'LEFT', 'RIGHT', 'MID', 'LEN',
                   'TRIM', 'UPPER', 'LOWER', 'VALUE', 'ISNUMBER', 'ISTEXT', 'ISBLANK', 'ISERROR',
                   'ISNA', 'DATE', 'YEAR', 'MONTH', 'DAY', 'DAYS', 'TRUE', 'FALSE'}


def analyze_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.formula.tokenizer import Tokenizer
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    path = Path(path)
    if path.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValueError('The pilot reads .xlsx and .xlsm workbooks')
    if path.stat().st_size > 32 * 1024**2:
        raise ValueError('Pilot workbook limit: 32 MB compressed')
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 256 * 1024**2:
            raise ValueError('Pilot workbook limit: 256 MB expanded')
    workbook = load_workbook(path, data_only=False, keep_links=False, keep_vba=False)
    formulas = {}
    cells = {}
    labels = {}
    def key(sheet, coordinate): return sheet + '!' + coordinate.replace('$', '').upper()
    for sheet in workbook:
        if sheet.max_row * sheet.max_column > 100000:
            raise ValueError('Pilot worksheet bounding range limit: 100,000 cells')
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None: continue
                node = key(sheet.title, cell.coordinate)
                cells[node] = cell.data_type
                labels[node] = node
                if cell.data_type == 'f':
                    formulas[node] = cell.value if isinstance(cell.value, str) else None
    if not formulas:
        raise ValueError('No formulas found. Hardcoded equal values alone cannot establish shared provenance.')
    if len(cells) > 20000: raise ValueError('Pilot populated-cell limit: 20,000')
    dependencies = {}
    issues = defaultdict(list)

    def reference(token, current_sheet, names_seen=frozenset()):
        if '!' in token:
            sheet_name, address = token.rsplit('!', 1)
            if sheet_name.startswith("'") and sheet_name.endswith("'"):
                sheet_name = sheet_name[1:-1].replace("''", "'")
        else:
            sheet_name, address = current_sheet, token
        if '[' in sheet_name or ']' in sheet_name:
            raise ValueError('External workbook reference: ' + token)
        if ':' in sheet_name:
            endpoints = sheet_name.split(':')
            by_name = {s.casefold(): i for i, s in enumerate(workbook.sheetnames)}
            if len(endpoints) != 2 or any(s.casefold() not in by_name for s in endpoints):
                raise ValueError('Unknown 3D worksheet endpoint: ' + token)
            first, last = (by_name[s.casefold()] for s in endpoints)
            if first > last or not DIRECT.fullmatch(address):
                raise ValueError('Unsupported 3D worksheet range: ' + token)
            results = set()
            for sheet in workbook.sheetnames[first:last + 1]:
                results.update(reference("'" + sheet.replace("'", "''") + "'!" + address, current_sheet))
                if len(results) > 10000:
                    raise ValueError('Reference expansion exceeds 10,000 cells: ' + token)
            return results
        sheet_name = next((s for s in workbook.sheetnames if s.casefold() == sheet_name.casefold()), sheet_name)
        if sheet_name not in workbook.sheetnames:
            raise ValueError('Unknown worksheet: ' + sheet_name)
        if DIRECT.fullmatch(address):
            low_col, low_row, high_col, high_row = range_boundaries(address)
            if high_col < low_col or high_row < low_row:
                raise ValueError('Reversed cell range: ' + token)
            if high_col > 16384 or high_row > 1048576:
                raise ValueError('Reference is outside Excel bounds')
            if (high_col - low_col + 1) * (high_row - low_row + 1) > 10000:
                raise ValueError('Reference expansion exceeds 10,000 cells: ' + token)
            return {key(sheet_name, f'{get_column_letter(col)}{row}')
                    for col in range(low_col, high_col + 1) for row in range(low_row, high_row + 1)}
        if '[' in address or '#' in address or '@' in address:
            raise ValueError('Table, spill or implicit-intersection reference: ' + token)
        if address.casefold() in names_seen:
            raise ValueError('Recursive defined name: ' + address)
        def find_name(mapping):
            return next((v for k, v in mapping.items() if k.casefold() == address.casefold()), None)
        name = find_name(workbook[sheet_name].defined_names) or find_name(workbook.defined_names)
        if name is None: raise ValueError('Unresolved name or reference: ' + token)
        # Names are accepted only when they resolve to static coordinates.
        if name.type != 'RANGE': raise ValueError('Non-range defined name: ' + address)
        results = set()
        for destination_sheet, destination in name.destinations:
            results.update(reference("'" + destination_sheet.replace("'", "''") + "'!" + destination,
                                     current_sheet, names_seen | {address.casefold()}))
        if not results: raise ValueError('Empty defined name: ' + address)
        return results

    for node, formula in formulas.items():
        refs = set()
        if formula is None:
            issues[node].append('Array/shared formula representation unsupported')
        else:
            try:
                for token in Tokenizer(formula).items:
                    if token.type == 'FUNC' and token.subtype == 'OPEN':
                        function = token.value[:-1].upper().replace('_XLFN.', '').replace('_XLWS.', '')
                        if function in DYNAMIC or function not in KNOWN_FUNCTIONS:
                            issues[node].append('Dynamic or unsupported function: ' + function)
                    elif token.type == 'OPERAND' and token.subtype == 'RANGE':
                        try: refs.update(reference(token.value, node.rsplit('!', 1)[0]))
                        except ValueError as error: issues[node].append(str(error))
                    elif token.type == 'OPERAND' and token.subtype == 'ERROR':
                        issues[node].append('Formula contains error: ' + token.value)
            except Exception as error:
                issues[node].append('Formula parsing failed: ' + type(error).__name__)
        for ref in refs:
            if cells.get(ref) == 'e': issues[node].append('References an error-valued cell: ' + ref)
        dependencies[node] = refs

    # Mark every node downstream of a cycle without misclassifying the cycle as an input.
    pending = set(formulas)
    order = []
    while pending:
        ready = sorted(node for node in pending if not (dependencies[node] & pending))
        if not ready: break
        order.extend(ready)
        pending.difference_update(ready)
    for node in pending: issues[node].append('Cycle or dependency downstream of a cycle')
    blocked = set(issues)
    for node in order:
        if dependencies[node] & blocked:
            blocked.add(node)
            issues[node].append('Depends on an unresolved formula')
    safe = set(formulas) - blocked
    safe_order = [node for node in order if node in safe]
    roots = {ref for node in safe for ref in dependencies[node] if ref not in formulas}
    if len(roots) + len(safe) > 20000: raise ValueError('Expanded lineage exceeds 20,000 nodes')
    if not safe:
        raise ValueError('No formula has a fully resolved static reference graph. ' + '; '.join(f'{n}: {issues[n][0]}' for n in list(issues)[:5]))
    used = {ref for node in safe for ref in dependencies[node] if ref in safe}
    targets = tuple(sorted(safe - used))
    source_sets = {root: frozenset((root,)) for root in roots}
    membership_count = len(roots)
    for node in safe_order:
        source_sets[node] = frozenset().union(*(source_sets[p] for p in dependencies[node]))
        membership_count += len(source_sets[node])
        if membership_count > 500000:
            raise ValueError('Pilot lineage limit: 500,000 transitive source memberships')
    groups = defaultdict(list)
    for node in safe:
        if source_sets[node]: groups[tuple(sorted(source_sets[node]))].append(node)
    same_origins = [{'root_cells': list(key_), 'formula_cells': sorted(nodes)}
                    for key_, nodes in groups.items() if len(nodes) > 1]
    same_origins.sort(key=lambda g: (-len(g['formula_cells']), g['formula_cells']))
    metadata = {'title': path.name, 'input_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'origin': 'Workbook formulas parsed automatically; workbook values were not recalculated.',
                'labels': labels, 'formula_text': {n: formulas[n] for n in sorted(safe)},
                'formula_count': len(formulas), 'resolved_formula_count': len(safe),
                'unresolved': [{'cell': n, 'reasons': sorted(set(reasons))} for n, reasons in sorted(issues.items())],
                'same_origin_groups': same_origins,
                'assumptions': ['All static references are treated as potential dependencies, including unused IF branches and algebraic cancellations.',
                                'Distinct root cells are not evidence of statistically independent sources.',
                                'External links, dynamic references, table references and cyclic paths are unresolved; no external file was fetched.',
                                'Switching a source off reports structural sensitivity, not a false numeric value, audit failure or opinion.']}
    problem = DependencyProblem('workbook_' + metadata['input_sha256'][:12], {r: r for r in sorted(roots)},
                                tuple((n, tuple(sorted(dependencies[n]))) for n in safe_order), targets,
                                'static_formula_lineage', metadata)
    workbook.close()
    return problem
