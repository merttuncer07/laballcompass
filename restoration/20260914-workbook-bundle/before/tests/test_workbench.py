from itertools import product
import json
from pathlib import Path
import tempfile
import unittest
from workbench.contracts import DependencyProblem
from workbench.methods import METHODS, R207Batch
from workbench.problems import evidence_demo, random_dag, choice_blocks, reference_truth
from workbench.report import analyze_problem
from workbench.workbooks import analyze_workbook
from workbench.cli import save_analysis


class DependencyTests(unittest.TestCase):
    def test_shared_source_collapses_derivative_identities(self):
        p = evidence_demo()
        self.assertEqual(len(p.bases), 5)
        self.assertEqual(len(p.sources), 3)
        expected = [True, True, True, False]
        scenarios = [set(), {'client_ledger'}, {'external_confirmation'}, {'client_ledger', 'external_confirmation'}]
        for method in METHODS.values():
            self.assertEqual([r['claim'] for r in method(p).solve(scenarios)], expected)

    def test_all_truth_assignments_against_separate_oracle(self):
        for seed in (21, 73, 910):
            p = random_dag(seed, nodes=30, source_count=5)
            scenarios = [{s for s, bit in zip(p.sources, bits) if bit} for bits in product((False, True), repeat=5)]
            expected = [reference_truth(p, s) for s in scenarios]
            for method in METHODS.values(): self.assertEqual(method(p).solve(scenarios), expected)

    def test_analytical_choice_blocks(self):
        p = choice_blocks(24)
        for method in METHODS.values():
            got = method(p).solve([{'a0'}, {'a0', 'b0'}, {'a0', 'b1'}, set(p.sources)])
            self.assertEqual([r['target'] for r in got], [True, False, True, False])
        self.assertLess(len(R207Batch(p).circuit.nodes), 100)

    def test_report_finds_shared_single_and_joint_cuts(self):
        p = DependencyProblem('cuts', {'a':'A','b':'B','c':'C'}, (('t',('a','b')),('t',('a','c'))), ('t',))
        result = analyze_problem(p)
        self.assertEqual(result['source_impacts']['A'], ['t'])
        self.assertEqual(result['source_impacts']['B'], [])
        self.assertEqual(result['observed_cuts_up_to_two']['t'], [['A'], ['B','C']])

    def test_axiom_and_no_scenarios(self):
        p = DependencyProblem('axiom', {}, (('t',()),), ('t',))
        for method in METHODS.values():
            self.assertEqual(method(p).solve([set()]), [{'t': True}])
            self.assertEqual(method(p).solve([]), [])

    def test_invalid_contracts_rejected(self):
        with self.assertRaises(ValueError): DependencyProblem('cycle', {}, (('a',('b',)),('b',('a',))), ('a',))
        with self.assertRaises(ValueError): DependencyProblem('unknown', {}, (('a',('missing',)),), ('a',))
        with self.assertRaises(ValueError): DependencyProblem('overlap', {'a':'A'}, (('a',()),), ('a',))
        with self.assertRaises(ValueError): R207Batch(evidence_demo()).solve([{'missing'}])


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root/'input.xlsx'
        self.wb = Workbook(); self.ws = self.wb.active; self.ws.title = 'Data'
    def tearDown(self): self.temp.cleanup()
    def parse(self):
        self.wb.save(self.path)
        return analyze_workbook(self.path)

    def test_ranges_and_copy_chain_trace_same_input(self):
        self.ws['A1']=10; self.ws['B1']='=A1'; self.ws['C1']='=B1';self.ws['D1']='=SUM(A1:A1)'
        p=self.parse()
        self.assertEqual(p.sources, ('Data!A1',))
        self.assertEqual(set(p.targets), {'Data!C1','Data!D1'})
        self.assertEqual(set(p.metadata['same_origin_groups'][0]['formula_cells']), {'Data!B1','Data!C1','Data!D1'})
        self.assertTrue(all(not v for v in R207Batch(p).solve([{'Data!A1'}])[0].values()))

    def test_identical_hardcodes_are_not_merged(self):
        self.ws['A1']=10; self.ws['A2']=10;self.ws['B1']='=A1';self.ws['B2']='=A2'
        p=self.parse();self.assertEqual(len(p.sources),2)
        self.assertEqual(R207Batch(p).solve([{'Data!A1'}])[0], {'Data!B1':False,'Data!B2':True})

    def test_quoted_sheet_names_names_and_case(self):
        from openpyxl.workbook.defined_name import DefinedName
        other=self.wb.create_sheet("O'Brien Data");other['B2']=8
        self.wb.defined_names.add(DefinedName('TotalInput',attr_text="'O''Brien Data'!$B$2"))
        self.ws['A1']="='o''brien data'!$B$2";self.ws['A2']='=SUM(TotalInput)'
        p=self.parse();self.assertEqual(p.sources,("O'Brien Data!B2",))

    def test_unsupported_and_downstream_are_reported(self):
        self.ws['A1']=10;self.ws['B1']='=A1';self.ws['C1']='=INDIRECT("A1")';self.ws['D1']='=C1+1'
        self.ws['E1']='=[1]Sheet1!A1';self.ws['F1']='=Table1[Amount]'
        p=self.parse();self.assertEqual(p.metadata['resolved_formula_count'],1)
        self.assertEqual({r['cell'] for r in p.metadata['unresolved']}, {'Data!C1','Data!D1','Data!E1','Data!F1'})

    def test_three_dimensional_ranges_follow_tab_order_including_middle(self):
        for name in ('Start tab', 'Z middle', 'End tab', 'Outside'):
            sheet = self.wb.create_sheet(name)
            sheet['A1'] = 1; sheet['B1'] = 2
        self.ws['A1'] = "=SUM('start tab:END TAB'!A1:B1)"
        p = self.parse()
        expected = {s+'!'+c for s in ('Start tab', 'Z middle', 'End tab') for c in ('A1','B1')}
        self.assertEqual(set(p.sources), expected)
        self.assertFalse(R207Batch(p).solve([{'Z middle!B1'}])[0]['Data!A1'])

    def test_invalid_ranges_never_become_dependency_free_formulas(self):
        self.ws['A1'] = 1; self.ws['B1'] = '=A1'
        self.ws['C1'] = '=SUM(B2:A1)'; self.ws['D1'] = '=SUM(Data:Missing!A1)'
        p = self.parse()
        self.assertEqual(p.metadata['resolved_formula_count'], 1)
        self.assertEqual({r['cell'] for r in p.metadata['unresolved']}, {'Data!C1','Data!D1'})

    def test_independent_apache_poi_three_dimensional_fixture(self):
        fixture = Path(__file__).resolve().parents[1] / 'examples/public-workbooks/FormulaSheetRange.xlsx'
        p = analyze_workbook(fixture)
        self.assertEqual(p.metadata['formula_count'], 2)
        self.assertEqual(p.metadata['unresolved'], [])
        rules = dict(p.rules)
        self.assertEqual(set(rules['test!D11']), {f'Sheet{i}!A11' for i in range(2,6)})
        self.assertEqual(set(rules['test!D12']), {f'Sheet{i}!{c}12' for i in range(2,6) for c in 'ABC'})

    def test_cycles_do_not_become_independent_sources(self):
        self.ws['A1']=3;self.ws['B1']='=A1';self.ws['C1']='=D1';self.ws['D1']='=C1';self.ws['E1']='=C1+1'
        p=self.parse();self.assertEqual(p.sources,('Data!A1',))
        self.assertEqual(len(p.metadata['unresolved']),3)

    def test_strings_do_not_create_fake_references(self):
        self.ws['A1']=1;self.ws['B1']='=IF(A1,"Z99","Q28")'
        p=self.parse();self.assertEqual(p.sources,('Data!A1',))
        self.assertEqual(p.mode,'static_formula_lineage')

    def test_conditional_branches_are_conservative(self):
        self.ws['A1']=True;self.ws['A2']=10;self.ws['A3']=20;self.ws['B1']='=IF(A1,A2,A3)'
        p=self.parse();self.assertEqual(set(p.sources),{'Data!A1','Data!A2','Data!A3'})
        self.assertIn('unused IF branches',p.metadata['assumptions'][0])

    def test_error_cell_and_downstream_excluded(self):
        self.ws['A1']=3;self.ws['B1']='=A1';self.ws['C1']='#N/A';self.ws['D1']='=C1'
        p=self.parse();self.assertEqual([r['cell'] for r in p.metadata['unresolved']],['Data!D1'])

    def test_empty_and_unresolved_only_workbooks_explain_failure(self):
        with self.assertRaisesRegex(ValueError,'No formulas'):self.parse()
        self.ws['A1']='=INDIRECT("B1")'
        with self.assertRaisesRegex(ValueError,'No formula has'):self.parse()

    def test_analysis_preserves_input_and_existing_output(self):
        self.ws['A1']=7;self.ws['B1']='=A1';p=self.parse();before=self.path.read_bytes()
        output=save_analysis(p,self.root/'report')
        self.assertEqual(before,self.path.read_bytes());self.assertTrue(output.exists())
        self.assertEqual(json.loads((output.parent/'problem.json').read_text())['mode'],'static_formula_lineage')
        with self.assertRaises(ValueError):save_analysis(p,output.parent)


if __name__ == '__main__': unittest.main()
