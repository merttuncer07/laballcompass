from pathlib import Path
import json
import tempfile
import unittest
from openpyxl import Workbook
from workbench.content import compare_content
from workbench.cli import save_workbook_analysis


class ContentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
    def tearDown(self):
        self.tmp.cleanup()
    def books(self, changed=False):
        left, right = Workbook(), Workbook()
        for r in range(1,5):
            for c in range(1,4):
                value = 100*r + 3*c
                left.active.cell(r,c,value)
                right.active.cell(r+6,c+2,value)
        if changed:right.active['D9']=999
        paths = [self.root/'ledger.xlsx',self.root/'support.xlsx']
        for p,b in zip(paths,(left,right)):b.save(p)
        return paths,[left,right]
    def test_shifted_block_and_changed_interior_are_located(self):
        paths,books=self.books(changed=True)
        result=compare_content(paths,books)
        self.assertEqual(len(result['blocks']),1)
        b=result['blocks'][0]
        self.assertEqual(b['left']['range'],'A1:C4')
        self.assertEqual(b['right']['range'],'C7:E10')
        self.assertEqual(b['matching_cells'],11)
        self.assertEqual([(c['left'],c['right']) for c in b['cells'] if not c['equal']],[('B3','D9')])
    def test_formula_free_files_produce_a_usable_report(self):
        paths,_=self.books()
        report=save_workbook_analysis(paths,self.root/'report')
        data=json.loads(report.with_name('content.json').read_text())
        self.assertTrue(report.exists())
        self.assertFalse(data['lineage_available'])
        self.assertEqual(data['blocks'][0]['matching_cells'],12)
        self.assertFalse(report.with_name('problem.json').exists())
    def test_matches_never_merge_formula_roots(self):
        paths,books=self.books()
        books[0].active['E1']='=SUM(A1:C4)'
        books[1].active['G1']='=SUM(C7:E10)'
        for p,b in zip(paths,books):b.save(p)
        report=save_workbook_analysis(paths,self.root/'report')
        data=json.loads(report.with_name('content.json').read_text())
        problem=json.loads(report.with_name('problem.json').read_text())
        self.assertEqual(len(set(problem['bases'].values())),24)
        self.assertEqual(len(data['blocks'][0]['formula_targets']),2)
        self.assertTrue(report.with_name('lineage.html').exists())
    def test_numeric_text_and_numbers_do_not_match(self):
        paths,books=self.books()
        for row in books[1].active:
            for cell in row:
                if cell.value is not None:cell.value=str(cell.value)
        self.assertEqual(compare_content(paths,books)['blocks'],[])
    def test_formulas_are_not_compared_as_values(self):
        paths,books=self.books()
        for row in books[1].active:
            for cell in row:
                if cell.value is not None:cell.value='='+str(cell.value)
        data=compare_content(paths,books)
        self.assertEqual(data['blocks'],[])
        self.assertEqual(data['excluded_cells']['formula'],12)
    def test_common_boilerplate_does_not_seed_a_block(self):
        paths,books=self.books()
        for b in books:
            for row in b.active:
                for cell in row:cell.value=0
        self.assertEqual(compare_content(paths,books)['blocks'],[])
    def test_coincidental_equal_blocks_are_not_claimed_as_provenance(self):
        paths,books=self.books()
        result=compare_content(paths,books)
        self.assertIn('not established',result['blocks'][0]['interpretation'])
        self.assertNotIn('source_groups',result)
    def test_output_is_preserved(self):
        paths,_=self.books()
        folder=self.root/'report';folder.mkdir();(folder/'keep.txt').write_text('original')
        with self.assertRaisesRegex(ValueError,'preserved'):
            save_workbook_analysis(paths,folder)
        self.assertEqual((folder/'keep.txt').read_text(),'original')

    def test_independently_authored_apache_workbook_shared_table(self):
        from workbench.workbooks import open_workbooks
        fixture = Path(__file__).resolve().parents[1]/'examples/public-workbooks/FormulaEvalTestData_Copy.xlsx'
        with open_workbooks([fixture]) as (paths,books):
            result=compare_content(paths,books)
        block=next(b for b in result['blocks'] if b['left']['sheet']=='EverythingTests' and b['right']['sheet']=='FinanceLibTests')
        self.assertEqual((block['left']['range'],block['right']['range']),('T7:X19','B2:F14'))
        self.assertEqual(block['matching_cells'],63)
        self.assertEqual([(c['left'],c['right'],c['right_value']) for c in block['cells'] if not c['equal']],
                         [('T12','B7','=2/3'),('T13','B8','=2/3')])


if __name__=='__main__':unittest.main()
