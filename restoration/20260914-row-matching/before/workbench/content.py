"""Locate shared literal cell blocks. Content matches never merge source identities.

This is a deterministic, bounded candidate finder, not a learned provenance
classifier. Typed rare-value anchors propose translations; actual cells verify
each block. No tolerance, formatting, filename or formula-cache similarity is
silently treated as value equality.
"""
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from itertools import combinations
import hashlib
import math


def typed_value(cell):
    value = cell.value
    if value is None or cell.data_type in ('f', 'e'):
        return None
    if isinstance(value, bool):
        return ('bool', value)
    if isinstance(value, (int, float)):
        return ('number', value) if math.isfinite(value) else None
    if isinstance(value, str):
        return ('text', value) if value.strip() else None
    if isinstance(value, (datetime, date, time)):
        return (type(value).__name__, value.isoformat())
    if isinstance(value, timedelta):
        return ('duration_seconds', value.total_seconds())
    return None


def _components(points):
    """Join matching cells across at most one intervening changed cell."""
    pending = set(points)
    offsets = [(r, c) for r in range(-2, 3) for c in range(-2, 3) if (r, c) != (0, 0)]
    while pending:
        start = min(pending)
        pending.remove(start)
        group, todo = {start}, [start]
        while todo:
            r, c = todo.pop()
            for dr, dc in offsets:
                point = (r + dr, c + dc)
                if point in pending:
                    pending.remove(point); group.add(point); todo.append(point)
        yield group


def compare_content(paths, books):
    from openpyxl.utils.cell import get_column_letter
    sheets, index, excluded = [], defaultdict(list), defaultdict(int)
    populated = 0
    for path, book in zip(paths, books):
        for sheet in book:
            if sheet.max_row * sheet.max_column > 100000:
                raise ValueError('Content comparison worksheet limit: 100,000 cells')
            values, raw = {}, {}
            sid = len(sheets)
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None: continue
                    populated += 1
                    if populated > 20000:
                        raise ValueError('Content comparison collection limit: 20,000 populated cells')
                    point = (cell.row, cell.column)
                    token = typed_value(cell)
                    raw[point] = str(cell.value)
                    if token is None:
                        excluded['formula' if cell.data_type == 'f' else 'error_or_empty_text'] += 1
                        continue
                    values[point] = token
                    # Trivial constants can verify a region, but cannot seed one.
                    if token[0] != 'bool' and token not in (('number', 0), ('number', 1), ('number', -1)):
                        index[token].append((sid, *point))
            sheets.append({'workbook': path.name, 'sheet': sheet.title, 'values': values, 'raw': raw})
    proposals = defaultdict(set)
    joins, common = 0, 0
    for token, positions in index.items():
        if len(positions) > 64:
            common += 1
            continue
        for (a, ar, ac), (b, br, bc) in combinations(positions, 2):
            if a == b: continue
            joins += 1
            if joins > 200000:
                raise ValueError('Content anchor comparison exceeds 200,000 pairs; supply a smaller collection')
            proposals[(a, b, br-ar, bc-ac)].add(token)
    candidates = sorted(k for k, tokens in proposals.items() if len(tokens) >= 3)
    if len(candidates) > 10000:
        raise ValueError('Too many content alignments; supply a smaller collection')
    comparisons, blocks = 0, []
    coord = lambda p: get_column_letter(p[1]) + str(p[0])
    for a, b, dr, dc in candidates:
        left, right = sheets[a], sheets[b]
        av, bv = left['values'], right['values']
        comparisons += len(av)
        if comparisons > 5000000:
            raise ValueError('Content comparison exceeds 5,000,000 cell checks; supply a smaller collection')
        matches = {p for p, token in av.items() if bv.get((p[0]+dr, p[1]+dc)) == token}
        for group in _components(matches):
            if len(group) < 6: continue
            if len({av[p] for p in group} & proposals[(a, b, dr, dc)]) < 3: continue
            r1, r2 = min(r for r,c in group), max(r for r,c in group)
            c1, c2 = min(c for r,c in group), max(c for r,c in group)
            area = (r2-r1+1)*(c2-c1+1)
            if len(group)/area < .8: continue
            blocks.append({
                'left': {'workbook':left['workbook'], 'sheet':left['sheet'], 'range':coord((r1,c1))+':'+coord((r2,c2))},
                'right': {'workbook':right['workbook'], 'sheet':right['sheet'], 'range':coord((r1+dr,c1+dc))+':'+coord((r2+dr,c2+dc))},
                'matching_cells':len(group), 'compared_positions':area, 'match_fraction':len(group)/area,
                'different_or_uncompared_cells':area-len(group),
                '_region':(a,b,r1,r2,c1,c2,dr,dc),
                'interpretation':'Observed content overlap; direction, copying and common origin are not established',
            })
    blocks.sort(key=lambda b:(-b['matching_cells'], b['left']['workbook'], b['left']['sheet'], b['left']['range'], b['right']['workbook'], b['right']['sheet'], b['right']['range']))
    # Materialize bounded report details only after ranking candidate regions.
    selected, detail_cells = [], 0
    for block in blocks:
        if len(selected) >= 200 or detail_cells + block['compared_positions'] > 100000:
            continue
        a,b,r1,r2,c1,c2,dr,dc = block.pop('_region')
        left,right = sheets[a],sheets[b]
        av,bv = left['values'],right['values']
        cells = []
        for r in range(r1,r2+1):
            for c in range(c1,c2+1):
                p,q = (r,c),(r+dr,c+dc)
                cells.append({'left':coord(p), 'right':coord(q),
                              'left_value':left['raw'].get(p), 'right_value':right['raw'].get(q),
                              'left_type':av[p][0] if p in av else 'uncompared',
                              'right_type':bv[q][0] if q in bv else 'uncompared',
                              'equal':p in av and av[p] == bv.get(q)})
        block['cells'] = cells
        block['matching_cells'] = sum(c['equal'] for c in cells)
        block['different_or_uncompared_cells'] = len(cells)-block['matching_cells']
        block['match_fraction'] = block['matching_cells']/len(cells)
        selected.append(block)
        detail_cells += len(cells)
    return {'schema_version':1, 'method':'typed_rare_anchor_translation_v1',
            'input_files':[{'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in paths],
            'sheet_count':len(sheets), 'populated_cells':populated,
            'blocks':selected, 'total_blocks_found':len(blocks), 'blocks_omitted':len(blocks)-len(selected),
            'excluded_cells':dict(excluded), 'common_anchor_values_skipped':common,
            'candidate_alignments':len(candidates), 'verified_cell_lookups':comparisons,
            'scope':[
                'Only supplied workbooks; comparison is between distinct sheets, including sheets within one workbook.',
                'At least 6 matching cells, 3 distinct nontrivial rare values, and 80% matching positions within the displayed rectangle.',
                'Exact typed equality only: numeric text differs from a number; formatting is ignored. Formula cells, caches and errors cannot establish a match.',
                'Only row/column translations are searched. Reordering, transposition, rounding and removed columns can be missed. Differences beyond the matching anchors are not located.',
                'No match means no match within this search, not independent evidence. Matched cells are never merged into one lineage root.',
                'Thresholds are review heuristics, not calibrated confidence or an audit conclusion.',
                'Output is limited to 200 regions and 100,000 detailed positions, ranked by matching cell count.',
            ]}
