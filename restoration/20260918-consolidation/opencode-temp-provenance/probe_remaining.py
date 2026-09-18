import json
import tempfile
from pathlib import Path
from openpyxl import Workbook
from workbench.content import compare_content
from workbench.record_matches import find_record_matches

with tempfile.TemporaryDirectory(dir='/var/folders/87/7zcwj9x95fq3g0bv00f1t_gh0000gn/T/opencode') as td:
    root = Path(td)
    book = Workbook()
    left = book.active
    left.title = 'Left'
    right = book.create_sheet('Right')
    for r, rr in enumerate([3,1,5,2,4], 1):
        for c in range(1,3):
            left.cell(r,c,1000*c+17*r)
            right.cell(rr,c,1000*c+17*r)
    left['C1'] = '=A1'
    left['D1'] = '#DIV/0!'
    left['A6'] = '=A1'
    left['A7'] = '#N/A'
    path = root/'unmapped.xlsx'
    book.save(path)
    result = compare_content([path], [book])
    block = next(b for b in result['blocks'] if b['kind']=='row_alignment')
    print(json.dumps({'probe':'unmapped','excluded':result['excluded_cells'],'actual':{k:v for k,v in block.items() if k.startswith('unmapped')},'expected_left_rows':[6,7],'expected_left_columns':['C','D']}))
    sheets = []
    for sid in range(33):
        values = {(r,c):('number',1000*c+r) for r in range(1,201 if sid>=31 else 101) for c in (1,2)}
        sheets.append({'workbook':'rank.xlsx','sheet':f'S{sid:02}','values':values,'raw':{p:str(t[1]) for p,t in values.items()}})
    for name, order in [('original',sheets),('strong_first',sheets[-2:]+sheets[:-2])]:
        blocks,scope = find_record_matches(order)
        strong = [b for b in blocks if {b['left']['sheet'],b['right']['sheet']}=={'S31','S32'}]
        print(json.dumps({'probe':'ranking','order':name,'blocks':len(blocks),'detail_cells':sum(len(b['cells']) for b in blocks),'max_matching':max(b['matching_cells'] for b in blocks),'strong_pair_present':bool(strong),'scope':scope}))
