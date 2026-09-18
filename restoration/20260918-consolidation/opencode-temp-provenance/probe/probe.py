import sys, time, resource
from pathlib import Path
sys.path.insert(0, '/Users/mertalituncer/Documents/Default Project/laballcompass-fixes')
from openpyxl import load_workbook

dirp = Path('/var/folders/87/7zcwj9x95fq3g0bv00f1t_gh0000gn/T/opencode/soni-pair-review')
for name in ('draft-model.xlsx', 'final-model.xlsx'):
    path = dirp / name
    t0 = time.time()
    wb = load_workbook(path, data_only=False, keep_links=True, keep_vba=False)
    populated = 0
    maxdim = 0
    for sheet in wb:
        dim = sheet.max_row * sheet.max_column
        maxdim = max(maxdim, dim)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    populated += 1
    print(f'{name}: populated={populated} max_sheet_cells={maxdim} load+scan={time.time()-t0:.1f}s')
    wb.close()
