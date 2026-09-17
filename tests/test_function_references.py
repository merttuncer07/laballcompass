"""Formula semantics counterexamples; not invented field/audit workbooks."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table
from workbench.workbooks import analyze_workbook, analyze_workbooks
from workbench.methods import METHODS


class ConditionalRangeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root/'formula-unit.xlsx'
        self.book = Workbook(); self.sheet = self.book.active; self.sheet.title = 'Data'
        for row in range(1, 11):
            for col in range(1, 4): self.sheet.cell(row, col, row)
        self.sheet['H1'] = '=1'  # Independent resolved formula exposes unresolved cases.

    def tearDown(self):
        self.book.close(); self.temp.cleanup()

    def parse(self, formula):
        self.sheet['G1'] = formula
        self.book.save(self.path)
        return analyze_workbook(self.path)

    def refs(self, problem):
        return set(dict(problem.rules)['Data!G1'])

    def test_single_cell_value_argument_expands_and_changes_source_impact(self):
        for function in ['SUMIF', 'AVERAGEIF']:
            with self.subTest(function=function):
                p = self.parse(f'={function}(A1:A3,">0",C1)')
                self.assertEqual(self.refs(p), {f'Data!{c}{r}' for c in ['A', 'C'] for r in range(1, 4)})
                self.assertEqual(p.metadata['function_range_adjustments'][0]['effective_reference'], 'Data!C1:C3')
                for method in METHODS.values():
                    self.assertFalse(method(p).solve([{'Data!C3'}])[0]['Data!G1'])

    def test_oversized_value_range_shrinks_without_spurious_error_dependency(self):
        self.sheet['C9'] = '#REF!'
        p = self.parse('=SUMIF(A1:A2,">0",C1:C10)')
        self.assertEqual(self.refs(p), {'Data!A1', 'Data!A2', 'Data!C1', 'Data!C2'})
        self.assertEqual(p.metadata['unresolved'], [])

    def test_two_dimensional_shape_extends_both_axes(self):
        p = self.parse('=AVERAGEIF(A1:B2,">0",D4)')
        self.assertTrue({'Data!D4', 'Data!E4', 'Data!D5', 'Data!E5'} <= self.refs(p))
        self.assertEqual(len(self.refs(p)), 8)

    def test_nested_calls_criteria_expression_and_whitespace(self):
        p = self.parse('=IF(SUMIF( ( A1:A3 ),IF(B1=1,">0", "x,y"), ( C1 ))>0,SUMIF(A1:A2,1,C1),0)')
        self.assertEqual(self.refs(p), {'Data!B1'} | {f'Data!{c}{r}' for c in ['A', 'C'] for r in range(1, 4)})
        self.assertEqual(len(p.metadata['function_range_adjustments']), 2)

    def test_override_is_local_to_its_argument_occurrence(self):
        p = self.parse('=SUMIF(A1:A2,">0",C1:C10)+SUM(C1:C10)')
        self.assertIn('Data!C10', self.refs(p))

    def test_equal_geometry_and_two_argument_forms_remain_unchanged(self):
        for formula in ['=SUMIF(A1:A3,">0",C1:C3)', '=SUMIF(A1:A3,">0")', '=AVERAGEIF(A1:A3,">0")']:
            with self.subTest(formula=formula):
                p = self.parse(formula)
                self.assertEqual(p.metadata['function_range_adjustments'], [])
                self.assertEqual(p.metadata['unresolved'], [])

    def test_static_defined_names_preserve_geometry(self):
        self.book.defined_names.add(DefinedName('Conditions', attr_text="'Data'!$A$1:$A$3"))
        self.book.defined_names.add(DefinedName('Start', attr_text="'Data'!$C$1"))
        self.assertIn('Data!C3', self.refs(self.parse('=SUMIF(Conditions,">0",Start)')))

    def test_multi_area_name_not_mistaken_for_adjacent_single_rectangle(self):
        self.book.defined_names.add(DefinedName('Joined', attr_text="'Data'!$A$1:$A$2,'Data'!$A$3:$A$4"))
        p = self.parse('=SUMIF(Joined,">0",C1)')
        self.assertNotIn('Data!G1', dict(p.rules))
        self.assertIn('Multi-area', str(p.metadata['unresolved']))

    def test_table_column_criteria_resizes_single_start_cell(self):
        self.sheet['A1']='Criterion'; self.sheet['B1']='Other'; self.sheet['C1']='Value'
        self.sheet.add_table(Table(displayName='Inputs', ref='A1:C4'))
        p = self.parse('=SUMIF(Inputs[Criterion],">0",C2)')
        self.assertEqual(self.refs(p), {f'Data!{c}{r}' for c in ['A', 'C'] for r in range(2, 5)})

    def test_out_of_bounds_or_calculated_ranges_are_explicitly_unresolved(self):
        for formula in ['=SUMIF(A1:B2,">0",XFD1)', '=SUMIF(A1:A2,">0",IF(B1,C1,C2))', '=SUMIF(A1:A2,1,C1,C2)']:
            with self.subTest(formula=formula):
                p = self.parse(formula)
                self.assertNotIn('Data!G1', dict(p.rules))
                self.assertTrue(p.metadata['unresolved'])

    def test_newly_exposed_error_blocks_downstream_formula(self):
        self.sheet['C3'] = '#REF!'; self.sheet['G2'] = '=G1+1'
        p = self.parse('=SUMIF(A1:A3,">0",C1)')
        self.assertEqual({i['cell'] for i in p.metadata['unresolved']}, {'Data!G1', 'Data!G2'})

    def test_supplied_workbook_value_range_is_resized_in_its_own_workbook(self):
        source = self.root/'source.xlsx'; self.book.save(source)
        other = Workbook(); sheet = other.active; sheet.title = 'Summary'
        sheet['A1'] = "=SUMIF('[source.xlsx]Data'!A1:A3,\">0\",'[source.xlsx]Data'!C1)"
        path = self.root/'summary.xlsx'; other.save(path); other.close()
        p = analyze_workbooks([source, path])
        refs = set(dict(p.rules)['summary.xlsx::Summary!A1'])
        self.assertIn('source.xlsx::Data!C3', refs)
        self.assertEqual(p.metadata['function_range_adjustments'][0]['effective_reference'], 'source.xlsx::Data!C1:C3')


if __name__ == '__main__': unittest.main()
