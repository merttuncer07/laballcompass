import json
from pathlib import Path

from openpyxl import Workbook

from workbench.cli import save_workbook_analysis
from workbench.revision_triage import build_revision_triage
from workbench.structural_drift import build_structural_drift


def _exact(*, removed=(), added=(), external=None):
    return {
        "method": "same_layout_v1",
        "input_files": [{"name": "before.xlsx"}, {"name": "after.xlsx"}],
        "unmatched_sheets": {"left": list(removed), "right": list(added)},
        "blocks": [],
        "external_dependencies": external or {
            "parser_available": True, "references": [], "limitations": []},
        "dependency_coverage": {"lineage_available": True,
                                "formula_count": 10, "resolved_formula_count": 8,
                                "unresolved_formula_count": 2},
    }


def _triage(groups=()):
    return {
        "groups": list(groups),
        "dependency_coverage": {"lineage_available": True,
                                "formula_count": 10, "resolved_formula_count": 8,
                                "unresolved_formula_count": 2},
    }


def _row_block(left="Old", right="New", *, header_rename=True):
    mapping = [{"left": "A", "right": "B"},
               {"left": "B", "right": "A"},
               {"left": "C", "right": "C"}]
    rows = [{"left": 1, "right": 1}, {"left": 2, "right": 3},
            {"left": 3, "right": 2}]
    cells = []
    values = {("A", "B"): (("Old amount", "New amount") if header_rename
                             else ("Amount", "Amount")),
              ("B", "A"): ("Account", "Account"),
              ("C", "C"): ("Date", "Date")}
    for pair in mapping:
        before, after = values[pair["left"], pair["right"]]
        cells.append({"left": pair["left"] + "1", "right": pair["right"] + "1",
                      "left_value": before, "right_value": after, "equal": before == after})
    return {
        "kind": "row_alignment", "left": {"sheet": left}, "right": {"sheet": right},
        "matching_rows": 3, "row_mapping": rows, "column_mapping": mapping,
        "unmapped_left_rows": [4], "unmapped_right_rows": [5, 6],
        "unmapped_left_columns": ["D"], "unmapped_right_columns": ["E"],
        "matching_cells": 8, "compared_positions": 9, "match_fraction": 0.95,
        "different_or_uncompared_cells": int(header_rename), "cells": cells,
    }


def test_sheet_added_removed_and_strong_rename_candidate():
    structural = {"blocks": [_row_block()], "record_alignment": {"ambiguous_rows_omitted": 0}}
    drift = build_structural_drift(
        _exact(removed=["Old"], added=["New"]), _triage(), structural)
    workbook = drift["workbook_structure"]
    assert workbook["sheets_added"] == ["New"]
    assert workbook["sheets_removed"] == ["Old"]
    assert [(row["left_sheet"], row["right_sheet"])
            for row in workbook["likely_sheet_renames"]] == [("Old", "New")]


def test_ambiguous_sheet_rename_is_withheld():
    first = _row_block("Old", "New A")
    second = _row_block("Old", "New B")
    structural = {"blocks": [first, second], "record_alignment": {"ambiguous_rows_omitted": 0}}
    workbook = build_structural_drift(
        _exact(removed=["Old"], added=["New A", "New B"]), _triage(), structural
    )["workbook_structure"]
    assert workbook["likely_sheet_renames"] == []
    assert len(workbook["rename_candidates_withheld"]) == 2


def test_column_record_drift_and_pure_reorder_are_compact():
    structural = {"blocks": [_row_block()], "record_alignment": {"ambiguous_rows_omitted": 2}}
    records = build_structural_drift(_exact(), _triage(), structural)["record_structure"]
    assert [row["column"] for row in records["columns_added"]] == ["E"]
    assert [row["column"] for row in records["columns_removed"]] == ["D"]
    assert len(records["columns_reordered"]) == 1
    assert records["likely_column_renames"][0]["before"] == "Old amount"
    assert records["regions"][0]["pure_row_reorder"] is False
    assert records["regions"][0]["matched_record_count"] == 3
    assert records["column_mapping_withheld"][-1]["reason"].startswith("2 ambiguous")

    pure = _row_block(header_rename=False)
    pure_records = build_structural_drift(
        _exact(), _triage(), {"blocks": [pure],
                              "record_alignment": {"ambiguous_rows_omitted": 0}}
    )["record_structure"]
    assert pure_records["regions"][0]["pure_row_reorder"] is True


def test_weak_column_mapping_is_withheld():
    block = _row_block(); block["match_fraction"] = 0.7
    records = build_structural_drift(
        _exact(), _triage(), {"blocks": [block],
                              "record_alignment": {"ambiguous_rows_omitted": 0}}
    )["record_structure"]
    assert records["columns_added"] == []
    assert records["columns_removed"] == []
    assert records["likely_column_renames"] == []
    assert records["column_mapping_withheld"]


def test_formula_transitions_link_to_existing_triage_groups():
    groups = []
    for index, (category, transition) in enumerate((
            ("formula_logic_change", None), ("formula_to_literal", None),
            ("literal_to_formula", None), ("added_content", "formula_introduced"),
            ("removed_content", "formula_removed"))):
        groups.append({"id": f"g{index}", "category": category,
                       "formula_transition": transition, "changed_cell_count": index + 1})
    formulas = build_structural_drift(_exact(), _triage(groups))["formula_structure"]
    assert formulas["formula_to_literal"]["group_count"] == 1
    assert formulas["literal_to_formula"]["group_count"] == 1
    assert formulas["formulas_introduced"]["triage_group_ids"] == ["g3"]
    assert formulas["formulas_removed"]["changed_cell_count"] == 5


def test_external_link_introduced_removed_and_missing_are_parser_observations(tmp_path):
    before, after = Workbook(), Workbook()
    for book, reference in ((before, "OldSource.xlsx"), (after, "NewSource.xlsx")):
        book.active["A1"] = "=1+1"
        book.active["A2"] = f"='[{reference}]Data'!A1"
    paths = [tmp_path / "before.xlsx", tmp_path / "after.xlsx"]
    before.save(paths[0]); after.save(paths[1])
    report = save_workbook_analysis(paths, tmp_path / "comparison", same_layout=True)
    exact = json.loads(report.with_name("content.json").read_text())
    triage = build_revision_triage(exact)
    external = build_structural_drift(exact, triage)["external_dependencies"]
    assert external["introduced"] == ["newsource.xlsx"]
    assert external["removed"] == ["oldsource.xlsx"]
    assert {row["reference"] for row in external["unresolved_or_missing"]} == {
        "OldSource.xlsx", "NewSource.xlsx"}
    assert external["resolved_against_supplied_evidence"] == []


def test_external_link_resolved_only_by_unique_supplied_filename(tmp_path):
    before, after = Workbook(), Workbook()
    before.active.title = "Data"; after.active.title = "Data"
    before.active["A1"] = "='[after.xlsx]Data'!A2"
    before.active["B1"] = "=1+1"
    after.active["A2"] = 42
    after.active["B1"] = "=1+1"
    paths = [tmp_path / "before.xlsx", tmp_path / "after.xlsx"]
    before.save(paths[0]); after.save(paths[1])
    report = save_workbook_analysis(paths, tmp_path / "resolved", same_layout=True)
    exact = json.loads(report.with_name("content.json").read_text())
    external = build_structural_drift(
        exact, build_revision_triage(exact))["external_dependencies"]
    assert len(external["resolved_against_supplied_evidence"]) == 1
    resolved = external["resolved_against_supplied_evidence"][0]
    assert resolved["from_workbook"] == "before.xlsx"
    assert resolved["matched_workbook"] == "after.xlsx"
    assert "version and authenticity" not in resolved["identity_basis"].lower()
