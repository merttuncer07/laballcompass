import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import unittest

from openpyxl import Workbook

from workbench.evidence_workspace import (
    confirm_candidate,
    create_workspace,
    import_folder,
    reject_candidate,
    workspace_state,
)


class EvidenceWorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.incoming = self.root / "incoming"
        self.incoming.mkdir()
        self.workspace = self.root / "workspace"
        create_workspace(self.workspace, "Revenue evidence")

    def workbook(self, path, *, changed_row=None, unrelated=False, reordered=False):
        book = Workbook()
        sheet = book.active
        sheet.title = "Unrelated" if unrelated else "Revenue"
        rows = list(range(1, 31))
        if reordered:
            rows = rows[10:] + rows[:10]
        for output_row, record in enumerate(rows, 2):
            if unrelated:
                values = [f"OTHER-{record}", f"Vendor {record}", record * 1007, f"Zone {record % 7}"]
            else:
                amount = record * 100
                if record == changed_row:
                    amount += 25
                values = [f"INV-{record:03}", f"Customer {record}", amount, f"Region {record % 4}"]
            for column, value in enumerate(values, 1):
                sheet.cell(output_row, column, value)
        for column, value in enumerate(("Invoice", "Customer", "Amount", "Region"), 1):
            sheet.cell(1, column, value)
        sheet["F2"] = "=SUM(C2:C31)"
        sheet["F3"] = "=AVERAGE(C2:C31)"
        sheet["F4"] = "=MAX(C2:C31)"
        book.save(path)
        book.close()

    def test_many_file_import_confirmation_persistence_and_idempotence(self):
        baseline = self.incoming / "schedule-original.xlsx"
        renamed_revision = self.incoming / "client-final-renamed.xlsx"
        duplicate = self.incoming / "copy-from-email.xlsx"
        unrelated = self.incoming / "other-area.xlsx"
        self.workbook(baseline)
        self.workbook(renamed_revision, changed_row=17)
        shutil.copyfile(baseline, duplicate)
        self.workbook(unrelated, unrelated=True)
        original_hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest()
                           for path in self.incoming.iterdir()}

        imported = import_folder(self.workspace, self.incoming)
        self.assertEqual(imported["new_versions"], 3)
        self.assertEqual(imported["exact_duplicates"], 1)
        state = workspace_state(self.workspace)
        self.assertEqual(len(state["versions"]), 3)
        pending = [candidate for candidate in state["candidates"] if candidate["status"] == "pending"]
        self.assertEqual(len(pending), 1)
        candidate = pending[0]
        names = {
            next(version for version in state["versions"] if version["id"] == version_id)["first_name"]
            for version_id in (candidate["left_version_id"], candidate["right_version_id"])
        }
        self.assertNotIn("other-area.xlsx", names)
        self.assertIn("Same named-sheet coordinates", candidate["evidence"]["reasons"][0])
        baseline_hash = original_hashes[baseline]
        before_id = next(version["id"] for version in state["versions"] if version["sha256"] == baseline_hash)

        confirmed = confirm_candidate(
            self.workspace, candidate["id"], before_id, artifact_name="Revenue Schedule"
        )
        self.assertTrue(Path(confirmed["comparison_report"]).is_file())
        repeated_confirmation = confirm_candidate(self.workspace, candidate["id"], before_id)
        self.assertTrue(repeated_confirmation["existing"])
        reopened = workspace_state(self.workspace)
        self.assertEqual(reopened["artifacts"][0]["name"], "Revenue Schedule")
        self.assertEqual(reopened["artifacts"][0]["version_order"][0], before_id)
        self.assertFalse(reopened["artifacts"][0]["ordering_ambiguous"])
        self.assertEqual(len(reopened["comparisons"]), 1)
        content = json.loads(Path(confirmed["comparison_report"]).with_name("content.json").read_text())
        self.assertEqual(sum(block["changed_cells"] for block in content["blocks"]), 1)
        self.assertTrue(content["lineage_available"])

        repeated = import_folder(self.workspace, self.incoming)
        self.assertEqual(repeated["new_versions"], 0)
        self.assertEqual(repeated["exact_duplicates"], 4)
        repeated_state = workspace_state(self.workspace)
        self.assertEqual(len(repeated_state["versions"]), 3)
        self.assertEqual(len(repeated_state["candidates"]), 3)
        self.assertEqual(len(repeated_state["comparisons"]), 1)
        self.assertEqual(original_hashes, {path: hashlib.sha256(path.read_bytes()).hexdigest()
                                           for path in self.incoming.iterdir()})

    def test_reordered_records_use_existing_general_matcher(self):
        first = self.incoming / "first.xlsx"
        second = self.incoming / "renamed.xlsx"
        self.workbook(first)
        self.workbook(second, changed_row=8, reordered=True)
        import_folder(self.workspace, self.incoming)
        state = workspace_state(self.workspace)
        candidate = next(value for value in state["candidates"] if value["status"] == "pending")
        self.assertEqual(candidate["evidence"]["general_correspondence"]["kind"], "row_alignment")
        before_id = next(version["id"] for version in state["versions"] if version["first_name"] == "first.xlsx")
        result = confirm_candidate(self.workspace, candidate["id"], before_id, artifact_name="Moved rows")
        comparison = workspace_state(self.workspace)["comparisons"][0]
        self.assertIsNotNone(comparison["structural_report_path"])
        self.assertTrue((self.workspace / comparison["structural_report_path"]).is_file())
        self.assertTrue(Path(result["comparison_report"]).is_file())

    def test_ambiguous_proposals_and_rejection_are_persistent(self):
        for index in range(3):
            self.workbook(self.incoming / f"copy-{index}.xlsx", changed_row=10 + index)
        import_folder(self.workspace, self.incoming)
        state = workspace_state(self.workspace)
        pending = [candidate for candidate in state["candidates"] if candidate["status"] == "pending"]
        self.assertEqual(len(pending), 3)
        page = (self.workspace / "index.html").read_text()
        self.assertIn("Ambiguous", page)
        reject_candidate(self.workspace, pending[0]["id"])
        reopened = workspace_state(self.workspace)
        rejected = next(candidate for candidate in reopened["candidates"] if candidate["id"] == pending[0]["id"])
        self.assertEqual(rejected["status"], "rejected")

    def test_workspace_create_is_safe_and_idempotent(self):
        second = create_workspace(self.workspace, "Ignored new name")
        self.assertFalse(second["created"])
        self.assertEqual(second["name"], "Revenue evidence")
        nonempty = self.root / "not-empty"
        nonempty.mkdir()
        (nonempty / "keep.txt").write_text("keep")
        with self.assertRaisesRegex(ValueError, "empty"):
            create_workspace(nonempty)
        self.assertEqual((nonempty / "keep.txt").read_text(), "keep")


if __name__ == "__main__":
    unittest.main()
