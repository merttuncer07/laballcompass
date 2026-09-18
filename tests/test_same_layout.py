"""Software counterexamples for version comparison; not audit field evidence."""
import contextlib
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from workbench.cli import main, save_workbook_analysis
from workbench.content import compare_content
from workbench.workbooks import open_workbooks


class SameLayoutTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.paths = [self.root/'before.xlsx', self.root/'after.xlsx']
        self.books = [Workbook(), Workbook()]
        for book in self.books:
            self.addCleanup(book.close)
            book.active['A1'] = 10
            book.active['B1'] = '=A1*2'
            book.active['C1'] = '=5'

    def save(self):
        for path, book in zip(self.paths, self.books):
            book.save(path)

    def report(self):
        self.save()
        path = save_workbook_analysis(self.paths, self.root/'report', same_layout=True)
        return json.loads(path.with_name('content.json').read_text())

    def test_identical_formula_text_does_not_create_change_impact(self):
        block = self.report()['blocks'][0]
        self.assertEqual(block['matching_formula_cells'], 2)
        self.assertEqual(block['changed_cells'], 0)
        self.assertEqual(block['different_or_uncompared_cells'], 0)
        self.assertEqual(block['difference_formula_targets'], [])
        self.assertEqual(block['uncompared_formula_targets'], [])
        formula = next(c for c in block['cells'] if c['left'] == 'B1')
        self.assertEqual(formula['comparison'], 'same_formula')
        self.assertFalse(formula['equal'])  # Text equality never asserts evaluated equality.

    def test_identical_array_formula_text_is_compared_without_cached_values(self):
        for book in self.books:
            book.active['D1'] = ArrayFormula(ref='D1', text='=SUM(A1:C1)')
        block = self.report()['blocks'][0]
        formula = next(c for c in block['cells'] if c['left'] == 'D1')
        self.assertEqual(formula['comparison'], 'same_formula')
        self.assertEqual(block['uncompared_cells'], 0)

    def test_changes_and_unknowns_have_distinct_counts_and_impacts(self):
        for book in self.books:
            book.active['D1'] = '=2'
            book.active['E1'] = '#N/A'
        self.books[0].active['F1'] = 123
        self.books[1].active['G1'] = 0
        self.books[1].active['A1'] = 11
        self.books[1].active['D1'] = '=3'
        block = self.report()['blocks'][0]
        states = {c['left']: c['comparison'] for c in block['cells']}
        self.assertEqual(states, {'A1':'changed_value', 'B1':'same_formula',
            'C1':'same_formula', 'D1':'changed_formula', 'E1':'uncompared',
            'F1':'removed', 'G1':'added'})
        self.assertEqual((block['changed_cells'], block['uncompared_cells']), (4, 1))
        self.assertEqual(block['difference_formula_targets'], [
            'after.xlsx::Sheet!B1', 'after.xlsx::Sheet!D1',
            'before.xlsx::Sheet!B1', 'before.xlsx::Sheet!D1'])

    def test_common_empty_and_unmatched_sheets_are_explicit(self):
        for book in self.books:
            book.create_sheet('Empty')
        self.books[0].create_sheet('Removed')['A1'] = 1
        self.books[1].create_sheet('Added')['A1'] = 2
        result = self.report()
        self.assertEqual(result['unmatched_sheets'], {'left':['Removed'], 'right':['Added']})
        empty = next(b for b in result['blocks'] if b['left']['sheet'] == 'Empty')
        self.assertEqual(empty['compared_positions'], 0)
        self.assertIsNone(empty['match_fraction'])

    def test_same_layout_preserves_loaded_snapshot_hashes(self):
        self.save()
        hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in self.paths]
        with open_workbooks(self.paths, preserve_order=True) as (paths, books):
            self.paths[0].write_bytes(b'replaced after loading')
            result = compare_content(paths, books, same_layout=True)
        self.assertEqual([f['sha256'] for f in result['input_files']], hashes)

    def test_cli_exposes_mode_and_accepts_over_20000_populated_cells(self):
        for book in self.books:
            for r in range(1, 1101):
                for c in range(1, 11):
                    book.active.cell(r, c, r*10+c)
            book.active['A1'] = '=1'
        self.save()
        output = self.root/'cli-report'
        with contextlib.redirect_stdout(io.StringIO()):
            code = main(['analyze', *map(str, self.paths), '--same-layout', '--output', str(output)])
        self.assertEqual(code, 0)
        result = json.loads((output/'content.json').read_text())
        self.assertEqual(result['populated_cells'], 22000)
        self.assertEqual(result['blocks'][0]['changed_cells'], 0)
        self.assertTrue(result['lineage_available'])

    def test_mode_requires_two_distinct_workbook_names(self):
        with self.assertRaisesRegex(ValueError, 'exactly two'):
            compare_content(self.paths[:1], self.books[:1], same_layout=True)
        with self.assertRaisesRegex(ValueError, 'distinct filenames'):
            compare_content([self.paths[0], self.paths[0]], self.books, same_layout=True)
