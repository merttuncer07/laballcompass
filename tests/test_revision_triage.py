import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from workbench.cli import save_workbook_analysis
from workbench.revision_triage import build_revision_triage, write_revision_triage


class RevisionTriageTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def save_pair(self, before, after, name="case"):
        paths = [self.root / f"{name}-before.xlsx", self.root / f"{name}-after.xlsx"]
        before.save(paths[0]); after.save(paths[1])
        before.close(); after.close()
        exact_report = save_workbook_analysis(paths, self.root / f"{name}-exact", same_layout=True)
        return paths, exact_report

    def test_formula_transitions_values_additions_and_impact_are_distinct(self):
        before, after = Workbook(), Workbook()
        values_before = {"A1": 10, "B1": "=A1*2", "C1": "=A1+3", "D1": 7,
                         "E1": "removed", "G1": "=A1+1", "H1": "=A1*10"}
        values_after = {"A1": 11, "B1": 22, "C1": "=A1+4", "D1": "=A1-4",
                        "F1": "added", "G1": "=A1+1", "I1": "=A1*20"}
        for coordinate, value in values_before.items(): before.active[coordinate] = value
        for coordinate, value in values_after.items(): after.active[coordinate] = value
        _, exact_report = self.save_pair(before, after, "transitions")
        destination = self.root / "triage"
        report = write_revision_triage(exact_report.with_name("content.json"), destination)
        result = json.loads((destination / "triage.json").read_text())
        exact = json.loads(exact_report.with_name("content.json").read_text())
        self.assertEqual(result, build_revision_triage(exact))

        self.assertEqual(result["summary"]["raw_change_count"], 8)
        self.assertEqual(result["group_type_distribution"], {
            "added_content": 2,
            "formula_logic_change": 1,
            "formula_to_literal": 1,
            "literal_to_formula": 1,
            "removed_content": 2,
            "value_change": 1,
        })
        self.assertEqual({group["formula_transition"] for group in result["groups"]
                          if group["formula_transition"]}, {"formula_introduced", "formula_removed"})
        self.assertGreater(next(group for group in result["groups"]
                                if group["category"] == "value_change")
                           ["resolved_downstream_target_count"], 0)
        self.assertIn('id="category"', report.read_text())
        self.assertIn('id="reach"', report.read_text())
        self.assertIn("../exact/index.html", report.read_text())

    def reordered_pair(self, *, changed=False, ambiguous=False):
        before, after = Workbook(), Workbook()
        order = [7, 3, 9, 1, 8, 6, 10, 5, 2, 4, 12, 11]
        for source_row in range(1, 13):
            target_row = order[source_row - 1]
            for column in range(1, 6):
                value = 17 * source_row if ambiguous else 1000 * column + 17 * source_row
                before.active.cell(source_row, column, value)
                after.active.cell(target_row, column, value)
        if changed:
            after.active.cell(order[3], 3, 999999)
        return before, after

    def triage_reorder(self, *, changed=False, ambiguous=False, name="reorder"):
        before, after = self.reordered_pair(changed=changed, ambiguous=ambiguous)
        paths, exact_report = self.save_pair(before, after, name)
        structural_report = save_workbook_analysis(paths, self.root / f"{name}-structural",
                                                   preserve_order=True)
        exact = json.loads(exact_report.with_name("content.json").read_text())
        structural = json.loads(structural_report.with_name("content.json").read_text())
        return build_revision_triage(exact, structural), structural

    def test_pure_row_reorder_is_normalized_but_raw_changes_remain_linked(self):
        result, structural = self.triage_reorder(name="pure")
        self.assertTrue(any(block["kind"] == "row_alignment" for block in structural["blocks"]))
        self.assertEqual(result["summary"]["movement_candidate_groups"], 1)
        self.assertGreater(result["summary"]["raw_change_count"], result["summary"]["triage_group_count"])
        movement = next(group for group in result["groups"] if group["category"] == "movement_candidate")
        self.assertGreater(movement["movement_evidence"]["raw_changed_positions_normalized"], 0)

    def test_mixed_reorder_preserves_actual_value_change(self):
        result, _ = self.triage_reorder(changed=True, name="mixed")
        self.assertEqual(result["summary"]["movement_candidate_groups"], 1)
        self.assertTrue(any(group["category"] == "value_change" for group in result["groups"]))
        observations = [observation for group in result["groups"]
                        if group["category"] == "value_change"
                        for observation in group["representative_observations"]]
        self.assertTrue(any("999999" in (observation["left_value"], observation["right_value"])
                            for observation in observations))

    def test_ambiguous_mapping_is_not_normalized(self):
        result, structural = self.triage_reorder(ambiguous=True, name="ambiguous")
        self.assertFalse(any(block["kind"] == "row_alignment" for block in structural["blocks"]))
        self.assertEqual(result["summary"]["movement_candidate_groups"], 0)
        self.assertGreater(result["summary"]["raw_change_count"], 0)


if __name__ == "__main__":
    unittest.main()
