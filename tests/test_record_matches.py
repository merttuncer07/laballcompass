import unittest
import tempfile
from pathlib import Path
from openpyxl import Workbook
from workbench.content import compare_content
from workbench.record_matches import find_record_matches


class RecordMatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def compare(self, paths, books):
        paths = [Path(self.tmp.name) / p.name for p in paths]
        for path, book in zip(paths, books):
            book.save(path)
        return compare_content(paths, books)

    def pair(self):
        a,b=Workbook(),Workbook()
        cols=[4,2,5,1,3]
        rows=[7,3,9,1,8,6,10,5,2,4]
        for r in range(1,11):
            for c in range(1,6):
                value=1000*c+17*r
                a.active.cell(r,c,value)
                b.active.cell(rows[r-1],cols[c-1],value)
        return [Path('one.xlsx'),Path('two.xlsx')],[a,b],rows,cols

    def test_late_stronger_region_is_ranked_before_bounded_details(self):
        reads=0
        class RawValues(dict):
            def get(self, key, default=None):
                nonlocal reads
                reads+=1
                return super().get(key,default)
        sheets=[]
        for sid in range(18):
            count=200 if sid<16 else 300
            offset=0 if sid<16 else 100000
            values={(r,c):('number',offset+10000*c+17*r)
                    for r in range(1,count+1) for c in range(1,6)}
            if sid==17:values[300,5]=('number',999999)
            sheets.append({'workbook':f'{sid:02}.xlsx','sheet':'Sheet','values':values,
                           'raw':RawValues({p:str(t[1]) for p,t in values.items()})})
        self.assertEqual(sum(len(s['raw']) for s in sheets),19000)
        blocks,scope=find_record_matches(sheets)
        self.assertEqual((blocks[0]['left']['workbook'],blocks[0]['right']['workbook']),('16.xlsx','17.xlsx'))
        self.assertEqual(blocks[0]['matching_cells'],1499)
        self.assertEqual([(c['left'],c['right']) for c in blocks[0]['cells'] if not c['equal']],[('E300','E300')])
        self.assertEqual(len(blocks),99)
        self.assertEqual(scope['regions_omitted_due_detail_limit'],22)
        self.assertEqual(sum(len(b['cells']) for b in blocks),99500)
        self.assertEqual(reads,199000)
        self.assertTrue(all(b['matching_cells']==sum(c['equal'] for c in b['cells']) for b in blocks))

    def test_unmapped_lists_include_raw_formula_and_error_only_coordinates(self):
        paths,books,rows,cols=self.pair()
        books[0].active['A12']='=SUM(A1:A10)'
        books[0].active['F1']='=A1'
        books[0].active['G14']='#N/A'
        books[1].active['B13']='#DIV/0!'
        books[1].active['H2']='#VALUE!'
        books[1].active['I15']='=SUM(A1:A10)'
        books[0].active['J20'].number_format='0.00'
        result=self.compare(paths,books)
        match=next(b for b in result['blocks'] if b.get('kind')=='row_alignment')
        self.assertEqual(match['unmapped_left_rows'],[12,14])
        self.assertEqual(match['unmapped_right_rows'],[13,15])
        self.assertEqual(match['unmapped_left_columns'],['F','G'])
        self.assertEqual(match['unmapped_right_columns'],['H','I'])
        self.assertEqual(match['matching_rows'],10)
        self.assertEqual(match['matching_cells'],50)
        self.assertEqual(match['compared_positions'],50)

    def test_reordered_rows_and_columns_locate_changed_value(self):
        paths,books,rows,cols=self.pair()
        books[1].active.cell(rows[3],cols[2],999999)
        result=self.compare(paths,books)
        match=next(b for b in result['blocks'] if b.get('kind')=='row_alignment')
        self.assertEqual(match['matching_rows'],10)
        self.assertEqual(match['matching_cells'],49)
        self.assertEqual(match['different_or_uncompared_cells'],1)
        self.assertEqual({(r['left'],r['right']) for r in match['row_mapping']},{(i+1,r) for i,r in enumerate(rows)})
        different=next(c for c in match['cells'] if not c['equal'])
        self.assertEqual((different['left'],different['right']),('C4','E1'))

    def test_ambiguous_identical_columns_are_not_arbitrarily_mapped(self):
        paths,books,rows,cols=self.pair()
        for book in books:
            for r in range(1,11):
                for c in range(1,6):book.active.cell(r,c,17*r)
        result=self.compare(paths,books)
        self.assertFalse(any(b.get('kind')=='row_alignment' for b in result['blocks']))

    def test_equal_column_bags_without_consistent_records_are_insufficient(self):
        a,b=Workbook(),Workbook()
        for r in range(12):
            for c in range(5):
                a.active.cell(r+1,c+1,1000*c+17*r)
                b.active.cell(r+1,c+1,1000*c+17*((r+c)%12))
        result=self.compare([Path('a.xlsx'),Path('b.xlsx')],[a,b])
        self.assertFalse(any(x.get('kind')=='row_alignment' for x in result['blocks']))

    def test_duplicate_rows_are_omitted_from_unique_correspondences(self):
        paths,books,rows,cols=self.pair()
        for c in range(1,6):
            books[0].active.cell(11,c,books[0].active.cell(1,c).value)
            books[1].active.cell(11,cols[c-1],books[0].active.cell(1,c).value)
        result=self.compare(paths,books)
        block=next(x for x in result['blocks'] if x.get('kind')=='row_alignment')
        self.assertFalse(any(x['left'] in (1,11) for x in block['row_mapping']))
        self.assertEqual(block['unmapped_left_rows'],[1,11])
        self.assertGreaterEqual(result['record_alignment']['ambiguous_rows_omitted'],2)

    def test_formula_difference_traces_only_resolved_downstream_targets(self):
        import json
        from workbench.cli import save_workbook_analysis
        for index,value in enumerate(('=C3+17','=3068','=INDIRECT("C3")','#N/A')):
            with self.subTest(value=value):
                paths,books,rows,cols=self.pair()
                books[0].active['C4']=value
                books[0].active['G1']='=SUM(C1:C10)'
                books[0].active['H1']='=G1+1'
                books[0].active['I1']='=C3'
                books[1].active['G1']='=SUM(E1:E10)'
                books[1].active['H1']='=G1+1'
                books[1].active['I1']="='[one.xlsx]Sheet'!C4"
                paths=[Path(self.tmp.name)/p.name for p in paths]
                for path,book in zip(paths,books):book.save(path)
                report=save_workbook_analysis(paths,Path(self.tmp.name)/f'report-{index}')
                content=json.loads(report.with_name('content.json').read_text())
                match=next(b for b in content['blocks'] if b.get('kind')=='row_alignment')
                expected=['one.xlsx::Sheet!H1','two.xlsx::Sheet!H1','two.xlsx::Sheet!I1'] if index<2 else ['two.xlsx::Sheet!H1']
                self.assertEqual(match['difference_formula_targets'],expected)
                different=[c for c in match['cells'] if not c['equal']]
                self.assertEqual([(c['left'],c['right']) for c in different],[('C4','E1')])
                self.assertEqual(different[0]['left_type'],'uncompared')
                self.assertIn('one.xlsx::Sheet!I1',match['formula_targets'])

    def test_existing_synthetic_changed_unique_voucher_pair(self):
        import hashlib
        import json
        from workbench.cli import save_workbook_analysis
        inputs = Path(__file__).resolve().parents[1] / 'examples/public-spending/verified/changed_unique_voucher/inputs'
        paths = [inputs / 'derived.xlsx', inputs / 'published.xlsx']
        before = [hashlib.sha256(path.read_bytes()).hexdigest() for path in paths]
        report = save_workbook_analysis(paths, Path(self.tmp.name) / 'existing-pair')
        content = json.loads(report.with_name('content.json').read_text())
        self.assertTrue(content['lineage_available'])
        self.assertEqual(len(content['blocks']), 1)
        block = content['blocks'][0]
        self.assertEqual(block['kind'], 'row_alignment')
        self.assertEqual(block['left']['workbook'], 'derived.xlsx')
        self.assertEqual(block['right']['workbook'], 'published.xlsx')
        self.assertEqual(block['matching_rows'], 85)
        self.assertEqual(block['matching_cells'], 679)
        self.assertEqual(block['compared_positions'], 680)
        self.assertEqual([cell for cell in block['cells'] if not cell['equal']], [{
            'left': 'C37', 'right': 'H8',
            'left_value': '42305.42', 'right_value': '42181.97',
            'left_type': 'number', 'right_type': 'number', 'equal': False,
        }])
        self.assertEqual(block['difference_formula_targets'], [
            'derived.xlsx::Review total!A2', 'published.xlsx::Review total!A2',
        ])
        self.assertEqual(content['blocks_omitted'], 0)
        self.assertEqual(before, [hashlib.sha256(path.read_bytes()).hexdigest() for path in paths])

    def test_reordered_difference_keeps_both_formula_impacts_and_separate_roots(self):
        import json
        from workbench.cli import save_workbook_analysis
        paths,books,rows,cols=self.pair()
        books[1].active.cell(rows[3],cols[2],999999)
        books[0].active['G1']='=SUM(C1:C10)'
        books[1].active['G1']='=SUM(E1:E10)'
        paths=[Path(self.tmp.name)/p.name for p in paths]
        for path,book in zip(paths,books):book.save(path)
        report=save_workbook_analysis(paths,Path(self.tmp.name)/'report')
        content=json.loads(report.with_name('content.json').read_text())
        match=next(b for b in content['blocks'] if b.get('kind')=='row_alignment')
        self.assertEqual(match['difference_formula_targets'],['one.xlsx::Sheet!G1','two.xlsx::Sheet!G1'])
        problem=json.loads(report.with_name('problem.json').read_text())
        self.assertEqual(len(set(problem['bases'].values())),20)


if __name__=='__main__':unittest.main()
