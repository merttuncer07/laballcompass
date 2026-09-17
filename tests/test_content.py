from pathlib import Path
import json
import tempfile
import unittest
from openpyxl import Workbook
from workbench.content import compare_content
from workbench.cli import save_workbook_analysis


class ContentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
    def tearDown(self):
        self.tmp.cleanup()
    def books(self, changed=False):
        left, right = Workbook(), Workbook()
        for r in range(1,5):
            for c in range(1,4):
                value = 100*r + 3*c
                left.active.cell(r,c,value)
                right.active.cell(r+6,c+2,value)
        if changed:right.active['D9']=999
        paths = [self.root/'ledger.xlsx',self.root/'support.xlsx']
        for p,b in zip(paths,(left,right)):b.save(p)
        return paths,[left,right]
    def test_same_layout_detects_changed_row_omitted_by_heuristics(self):
        paths,books=self.books()
        for book in books:
            book.remove(book.active)
            sheet=book.create_sheet('Data')
            for r in range(1,5):
                for c in range(1,11):sheet.cell(r,c,100*r+c)
        for c in range(1,11):books[1].active.cell(4,c,900+c)
        default=compare_content(paths,books)
        self.assertFalse(any(c['left']=='J4' for b in default['blocks'] for c in b['cells']))
        result=compare_content(paths,books,same_layout=True)
        block=result['blocks'][0]
        self.assertEqual(block['kind'],'same_layout')
        self.assertEqual(block['compared_positions'],40)
        self.assertEqual(block['matching_cells'],30)
        changed=next(c for c in block['cells'] if c['left']=='J4')
        self.assertEqual((changed['right'],changed['left_value'],changed['right_value'],changed['equal']),
                         ('J4','410','910',False))
        self.assertEqual(compare_content(paths,books,same_layout=False),default)

    def test_shifted_block_and_changed_interior_are_located(self):
        paths,books=self.books(changed=True)
        result=compare_content(paths,books)
        self.assertEqual(len(result['blocks']),1)
        b=result['blocks'][0]
        self.assertEqual(b['left']['range'],'A1:C4')
        self.assertEqual(b['right']['range'],'C7:E10')
        self.assertEqual(b['matching_cells'],11)
        self.assertEqual([(c['left'],c['right']) for c in b['cells'] if not c['equal']],[('B3','D9')])
    def test_formula_free_files_produce_a_usable_report(self):
        paths,_=self.books()
        report=save_workbook_analysis(paths,self.root/'report')
        data=json.loads(report.with_name('content.json').read_text())
        self.assertTrue(report.exists())
        self.assertFalse(data['lineage_available'])
        self.assertEqual(data['blocks'][0]['matching_cells'],12)
        self.assertFalse(report.with_name('problem.json').exists())
    def test_matches_never_merge_formula_roots(self):
        paths,books=self.books()
        books[0].active['E1']='=SUM(A1:C4)'
        books[1].active['G1']='=SUM(C7:E10)'
        for p,b in zip(paths,books):b.save(p)
        report=save_workbook_analysis(paths,self.root/'report')
        data=json.loads(report.with_name('content.json').read_text())
        problem=json.loads(report.with_name('problem.json').read_text())
        self.assertEqual(len(set(problem['bases'].values())),24)
        self.assertEqual(len(data['blocks'][0]['formula_targets']),2)
        self.assertTrue(report.with_name('lineage.html').exists())
    def test_formula_target_itself_is_included_for_single_workbook_difference(self):
        paths,books=self.books()
        books[0].active['B3']='=304+2'
        books[0].active['G1']='=A1'
        books[0].active['H1']='=INDIRECT("B3")'
        copied=books[0].create_sheet('Copy')
        for row in books[1].active:
            for cell in row:
                if cell.value is not None:copied.cell(cell.row,cell.column,cell.value)
        books[0].save(paths[0])
        report=save_workbook_analysis(paths[:1],self.root/'report')
        data=json.loads(report.with_name('content.json').read_text())
        block=next(b for b in data['blocks'] if b['kind']=='translated_block')
        self.assertEqual(block['difference_formula_targets'],['Sheet!B3'])
        self.assertEqual(block['formula_targets'],['Sheet!G1'])
        self.assertEqual([(c['left'],c['right']) for c in block['cells'] if not c['equal']],[('B3','D9')])

    def test_numeric_text_and_numbers_do_not_match(self):
        paths,books=self.books()
        for row in books[1].active:
            for cell in row:
                if cell.value is not None:cell.value=str(cell.value)
        self.assertEqual(compare_content(paths,books)['blocks'],[])
    def test_formulas_are_not_compared_as_values(self):
        paths,books=self.books()
        for row in books[1].active:
            for cell in row:
                if cell.value is not None:cell.value='='+str(cell.value)
        data=compare_content(paths,books)
        self.assertEqual(data['blocks'],[])
        self.assertEqual(data['excluded_cells']['formula'],12)
    def test_common_boilerplate_does_not_seed_a_block(self):
        paths,books=self.books()
        for b in books:
            for row in b.active:
                for cell in row:cell.value=0
        self.assertEqual(compare_content(paths,books)['blocks'],[])
    def test_coincidental_equal_blocks_are_not_claimed_as_provenance(self):
        paths,books=self.books()
        result=compare_content(paths,books)
        self.assertIn('not established',result['blocks'][0]['interpretation'])
        self.assertNotIn('source_groups',result)
    def test_output_is_preserved(self):
        paths,_=self.books()
        folder=self.root/'report';folder.mkdir();(folder/'keep.txt').write_text('original')
        with self.assertRaisesRegex(ValueError,'preserved'):
            save_workbook_analysis(paths,folder)
        self.assertEqual((folder/'keep.txt').read_text(),'original')

    def test_independently_authored_apache_workbook_shared_table(self):
        from workbench.workbooks import open_workbooks
        fixture = Path(__file__).resolve().parents[1]/'examples/public-workbooks/FormulaEvalTestData_Copy.xlsx'
        with open_workbooks([fixture]) as (paths,books):
            result=compare_content(paths,books)
        block=next(b for b in result['blocks'] if b['left']['sheet']=='EverythingTests' and b['right']['sheet']=='FinanceLibTests')
        self.assertEqual((block['left']['range'],block['right']['range']),('T7:X19','B2:F14'))
        self.assertEqual(block['matching_cells'],63)
        self.assertEqual([(c['left'],c['right'],c['right_value']) for c in block['cells'] if not c['equal']],
                         [('T12','B7','=2/3'),('T13','B8','=2/3')])


if __name__=='__main__':unittest.main()
