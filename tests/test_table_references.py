from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from openpyxl.worksheet.table import Table,TableColumn
from workbench.table_references import TableReferences,split_sheet_reference
from workbench.workbooks import analyze_workbook,analyze_workbooks


def table_book(headers=True,totals=True,names=('Amount','Rate','Computed')):
    book=Workbook();sheet=book.active;sheet.title='Data'
    sheet.append(list(names))
    for row in ((10,.1,1),(20,.2,4),(30,.3,9),(40,.4,16)):sheet.append(row)
    table=Table(displayName='Sales',ref='A1:C5',headerRowCount=int(headers),totalsRowCount=int(totals),
                tableColumns=[TableColumn(id=i+1,name=n) for i,n in enumerate(names)])
    sheet.add_table(table)
    return book


@pytest.mark.parametrize('token,expected',[
    ('Sales[Amount]',{'A2','A3','A4'}),
    ('Sales',{'A2','A3','A4','B2','B3','B4','C2','C3','C4'}),
    ('Sales[]',{'A2','A3','A4','B2','B3','B4','C2','C3','C4'}),
    ('Sales[[#Totals],[Amount]]',{'A5'}),
    ('Sales[[#Headers],[Amount]:[Rate]]',{'A1','B1'}),
    ('Sales[[#Data],[#Totals],[Rate]]',{'B2','B3','B4','B5'}),
    ('Sales[[#Headers],[#Data],[Rate]]',{'B1','B2','B3','B4'}),
    ('Sales[[#All],[Rate]]',{'B1','B2','B3','B4','B5'}),
    ('Sales[[#This Row],[Amount]]',{'A3'}),
    ('Sales[@Amount]',{'A3'}),
    ('Sales[@[Amount]:[Rate]]',{'A3','B3'}),
    ('[@Rate]',{'B3'}),
    ('[[#This Row],[Rate]]',{'B3'}),
    ('Sales[ [Amount]:[Rate] ]',{'A2','A3','A4','B2','B3','B4'}),
])
def test_selectors_resolve_exact_metadata_cells(token,expected):
    assert TableReferences([table_book()]).resolve(token,0,'Data','C3')==('Data',expected)


def test_escaped_header_characters_and_sheet_separator():
    name="x]#@!'"
    tables=TableReferences([table_book(names=(name,'Rate','Computed'))])
    token="Sales[x']'#'@!'']"
    assert split_sheet_reference(token)==(None,token)
    assert tables.resolve(token,0,'Data','C3')==('Data',{'A2','A3','A4'})
    assert split_sheet_reference("'O''Brien'!"+token)==("'O''Brien'",token)
    punctuation=TableReferences([table_book(names=(',',':','!'))])
    assert punctuation.resolve('Sales[[,]:[:]]',0,'Data','C3')[1]=={'A2','A3','A4','B2','B3','B4'}


def test_missing_headers_totals_and_outside_this_row_are_unresolved():
    tables=TableReferences([table_book(headers=False,totals=False)])
    assert tables.resolve('Sales[Amount]',0,'Data','C3')[1]=={'A1','A2','A3','A4','A5'}
    for token in ('Sales[#Headers]','Sales[#Totals]'):
        with pytest.raises(ValueError,match='no .* row'):tables.resolve(token,0,'Data','C3')
    tables=TableReferences([table_book()])
    for coordinate in ('C1','C5','C6'):
        with pytest.raises(ValueError,match='outside table data'):tables.resolve('Sales[@Amount]',0,'Data',coordinate)


@pytest.mark.parametrize('token',['Sales[[#All],[#This Row],[Amount]]','Sales[[Amount],[Rate]]','Sales[[#Headers],[#Totals]]',
                                'Sales[Unknown]','Sales[[Amount]','Sales[Amount]junk','Sales[[#Data],[Amount]:]','Sales[[]]'])
def test_malformed_or_unknown_reference_does_not_invent_roots(token):
    with pytest.raises(ValueError):TableReferences([table_book()]).resolve(token,0,'Data','C3')


def test_unqualified_reference_needs_containing_table():
    with pytest.raises(ValueError,match='outside'):TableReferences([table_book()]).resolve('[@Amount]',0,'Data','F3')


def test_same_table_name_in_different_workbooks_stays_separate(tmp_path):
    for name in ('a.xlsx','b.xlsx'):
        book=table_book();book['Data']['C3']='=[@Amount]*[@Rate]';book.save(tmp_path/name)
    p=analyze_workbooks([tmp_path/'a.xlsx',tmp_path/'b.xlsx'])
    assert set(p.bases)=={'a.xlsx::Data!A3','a.xlsx::Data!B3','b.xlsx::Data!A3','b.xlsx::Data!B3'}
    assert dict(p.rules)['a.xlsx::Data!C3']==('a.xlsx::Data!A3','a.xlsx::Data!B3')


def test_failed_table_reference_blocks_downstream_without_hiding_valid_formula(tmp_path):
    book=table_book(totals=False);s=book['Data']
    s['F1']='=SUM(Sales[#Totals])';s['F2']='=F1+1';s['C3']='=[@Amount]*[@Rate]'
    path=tmp_path/'table.xlsx';book.save(path);p=analyze_workbook(path)
    assert p.metadata['resolved_formula_count']==1
    assert {r['cell'] for r in p.metadata['unresolved']}=={'Data!F1','Data!F2'}
    assert set(p.bases)=={'Data!A3','Data!B3'}


def test_public_apache_references_match_independently_authored_expected_ranges():
    root=Path(__file__).resolve().parents[1]
    path=root/'examples/public-tables/StructuredReferences.xlsx'
    book=load_workbook(path);tables=TableReferences([book])
    # Expected values are from Apache POI TestXSSFFormulaParser.java, not this parser.
    for token,expected in [('\\_Prime.1[Name]',{'B2','B3','B4','B5','B6','B7'}),
                           ("\\_Prime.1[calc='#*'#]",{'A2','A3','A4','A5','A6','A7'}),
                           ('\\_Prime.1[[#Headers],[Number]]',{'C1'}),
                           ('\\_Prime.1[[#This Row],[Number]]',{'C3'})]:
        assert tables.resolve(token,0,'Table','C3')==('Table',expected)
    book.close()
    p=analyze_workbook(path)
    assert p.metadata['resolved_formula_count']==p.metadata['formula_count']==9
    assert dict(p.rules)['Table!A3']==('Table!C3',)
    assert set(dict(p.rules)['Formulas!A1'])=={'Table!A2','Table!A3','Table!A4','Table!A5','Table!A6','Table!A7'}


def test_public_apache_table_range_completes_formula_lineage():
    path=Path(__file__).resolve().parents[1]/'examples/public-tables/evaluate_formula_with_structured_table_references.xlsx'
    p=analyze_workbook(path)
    assert set(p.bases)=={'Tabelle1!A2','Tabelle1!A3','Tabelle1!B2','Tabelle1!B3'}
    assert dict(p.rules)['Tabelle1!C3']==tuple(sorted(p.bases))
