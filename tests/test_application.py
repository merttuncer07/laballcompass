import hashlib
from pathlib import Path

from openpyxl import Workbook

from workbench.application import EvidenceApplication


def _workbook(path, changed=False, formula=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Revenue"
    sheet.append(["Account", "Amount", "Calculated"])
    for row in range(2, 12):
        amount = row * 100 + (1 if changed and row == 5 else 0)
        sheet.append([f"A-{row}", amount, f"=B{row}*2"])
    if formula:
        sheet["C5"] = 1001 if changed else "=B5*2"
    book.save(path)


def _imported_application(tmp_path, formula=False):
    evidence = tmp_path / "evidence"
    evidence.mkdir()
    _workbook(evidence / "received.xlsx", formula=formula)
    _workbook(evidence / "revised.xlsx", changed=True, formula=formula)
    hashes = {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
              for path in evidence.iterdir()}
    application = EvidenceApplication()
    application.create_workspace(tmp_path / "workspace", "Revenue audit")
    result = application.import_folder(evidence)
    return application, evidence, hashes, result


def test_workspace_create_open_import_and_originals_unchanged(tmp_path):
    application, evidence, hashes, result = _imported_application(tmp_path)
    assert result["operation"]["new_blobs"] == 2
    assert len(application.list_unassigned_evidence()) == 2
    assert len(application.list_candidates(status="pending")) == 1
    assert hashes == {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                      for path in evidence.iterdir()}

    reopened = EvidenceApplication(tmp_path / "workspace")
    assert reopened.snapshot()["workspace"]["name"] == "Revenue audit"
    assert len(reopened.list_candidates(status="pending")) == 1


def test_candidate_rejection_persists(tmp_path):
    application, _, _, _ = _imported_application(tmp_path)
    candidate = application.list_candidates(status="pending")[0]
    application.reject_candidate(candidate["id"])
    reopened = EvidenceApplication(tmp_path / "workspace")
    assert reopened.list_candidates()[0]["status"] == "rejected"


def test_confirmed_history_and_triage_are_available_through_service(tmp_path):
    application, _, _, _ = _imported_application(tmp_path, formula=True)
    candidate = application.list_candidates(status="pending")[0]
    before_blob_id = (candidate["left_blob_id"] if candidate["left_name"] == "received.xlsx"
                      else candidate["right_blob_id"])
    confirmed = application.confirm_candidate(
        candidate["id"], before_blob_id, artifact_name="Revenue Schedule")
    relationship_id = confirmed["operation"]["relationship_id"]

    versions = application.list_versions(confirmed["operation"]["artifact_id"])
    assert [row["version_number"] for row in versions] == [1, 2]
    triage = application.get_revision_triage(relationship_id)
    assert triage["summary"]["triage_group_count"] >= 1
    assert "dependency_coverage" in triage
    assert "formula_to_literal" in {row["category"] for row in triage["groups"]}
    group = application.get_triage_group(relationship_id, triage["groups"][0]["id"])
    assert group["id"] == triage["groups"][0]["id"]


def test_correction_operations_flow_through_service(tmp_path):
    application, _, _, _ = _imported_application(tmp_path)
    candidate = application.list_candidates(status="pending")[0]
    result = application.confirm_candidate(
        candidate["id"], candidate["left_blob_id"], artifact_name="Old name")
    artifact_id = result["operation"]["artifact_id"]
    relationship_id = result["operation"]["relationship_id"]
    application.rename_artifact(artifact_id, "Revenue", "Corrected naming")
    relationship = next(row for row in application.snapshot()["relationships"]
                        if row["id"] == relationship_id)
    reversed_result = application.correct_version_order(
        relationship_id, relationship["after_version_id"], "Client clarified chronology")
    application.withdraw_relationship(
        reversed_result["operation"]["relationship_id"], "Relationship was mistaken")
    state = EvidenceApplication(tmp_path / "workspace").snapshot()
    assert state["artifacts"][0]["name"] == "Revenue"
    assert not any(row["status"] == "active" for row in state["relationships"])


def _candidate_for(application, first_name, second_name):
    wanted = {first_name, second_name}
    return next(row for row in application.list_candidates(status="pending")
                if {row["left_name"], row["right_name"]} == wanted)


def _confirm_named_pair(application, first_name, second_name, *, artifact_id=None,
                        artifact_name=None):
    candidate = _candidate_for(application, first_name, second_name)
    before_id = (candidate["left_blob_id"] if candidate["left_name"] == first_name
                 else candidate["right_blob_id"])
    return application.confirm_candidate(
        candidate["id"], before_id, artifact_id=artifact_id,
        artifact_name=artifact_name)


def test_use_checkpoint_identity_multiple_restart_and_changed_state(tmp_path):
    application, evidence, _, _ = _imported_application(tmp_path)
    confirmed = _confirm_named_pair(
        application, "received.xlsx", "revised.xlsx", artifact_name="Revenue")
    artifact_id = confirmed["operation"]["artifact_id"]
    versions = application.list_versions(artifact_id)
    used = versions[1]
    first = application.create_use_checkpoint(
        used["id"], "REV-03 testing", "Prepared from client schedule")
    second = application.create_use_checkpoint(used["id"], "Senior review 18 Sep")
    assert first["operation"]["version_id"] == used["id"]
    assert first["operation"]["blob_id"] == used["blob_id"]
    assert len(application.list_use_checkpoints()) == 2
    assert {row["change_state"] for row in application.list_use_checkpoints()} == {"CURRENT"}

    book = Workbook()
    sheet = book.active; sheet.title = "Revenue"
    sheet.append(["Account", "Amount", "Calculated"])
    for row in range(2, 12):
        sheet.append([f"A-{row}", row * 100 + (2 if row == 5 else 0), f"=B{row}*2"])
    book.save(evidence / "third.xlsx")
    application.import_folder(evidence)
    _confirm_named_pair(
        application, "revised.xlsx", "third.xlsx", artifact_id=artifact_id)

    changed = application.list_changed_since_use()
    assert len(changed) == 2
    assert all(row["version_id"] == used["id"] for row in changed)
    assert all(row["latest_version_number"] == 3 for row in changed)
    reopened = EvidenceApplication(tmp_path / "workspace")
    assert len(reopened.list_changed_since_use()) == 2
    assert reopened.get_changed_since_use(second["operation"]["checkpoint_id"],
                                           include_analysis=False)["version_id"] == used["id"]


def test_v1_checkpoint_compares_directly_to_v3_and_preserves_partial_coverage(tmp_path):
    application, evidence, _, _ = _imported_application(tmp_path, formula=True)
    confirmed = _confirm_named_pair(
        application, "received.xlsx", "revised.xlsx", artifact_name="Revenue")
    artifact_id = confirmed["operation"]["artifact_id"]
    first_version = application.list_versions(artifact_id)[0]
    checkpoint = application.create_use_checkpoint(
        first_version["id"], "Population prepared")["operation"]

    book = Workbook(); sheet = book.active; sheet.title = "Revenue"
    sheet.append(["Account", "Amount", "Calculated"])
    for row in range(2, 12):
        sheet.append([f"A-{row}", row * 100 + (2 if row == 5 else 0), f"=B{row}*2"])
    sheet["C5"] = 1002
    book.save(evidence / "third.xlsx")
    application.import_folder(evidence)
    _confirm_named_pair(application, "revised.xlsx", "third.xlsx", artifact_id=artifact_id)

    result = application.get_changed_since_use(checkpoint["checkpoint_id"])
    assert result["version_id"] == first_version["id"]
    assert result["latest_version_number"] == 3
    assert result["triage"]["comparison"]["before_version_id"] == first_version["id"]
    assert result["triage"]["comparison"]["after_version_id"] == result["latest_version_id"]
    assert result["triage"]["summary"]["triage_group_count"] >= 1
    assert "analysis_partial" in result
    cached = EvidenceApplication(tmp_path / "workspace").get_changed_since_use(
        checkpoint["checkpoint_id"])
    assert cached["triage"]["comparison"]["id"] == result["triage"]["comparison"]["id"]


def test_checkpoint_corrections_history_and_duplicate_do_not_create_change(tmp_path):
    application, evidence, _, _ = _imported_application(tmp_path)
    confirmed = _confirm_named_pair(
        application, "received.xlsx", "revised.xlsx", artifact_name="Revenue")
    latest = application.list_versions(confirmed["operation"]["artifact_id"])[-1]
    checkpoint = application.create_use_checkpoint(
        latest["id"], "REV-03", "Original note")["operation"]
    duplicate = tmp_path / "duplicate"
    duplicate.mkdir()
    (duplicate / "renamed-copy.xlsx").write_bytes((evidence / "revised.xlsx").read_bytes())
    application.import_folder(duplicate)
    assert application.get_changed_since_use(
        checkpoint["checkpoint_id"], include_analysis=False)["change_state"] == "CURRENT"

    application.update_checkpoint_metadata(
        checkpoint["checkpoint_id"], purpose="REV-03 testing", note="Corrected note",
        reason="Clarified external reference")
    updated = application.list_use_checkpoints()[0]
    assert updated["version_id"] == latest["id"]
    assert updated["purpose"] == "REV-03 testing"
    application.withdraw_use_checkpoint(checkpoint["checkpoint_id"], "Created accidentally")
    assert application.list_use_checkpoints() == []
    snapshot = application.snapshot()
    events = [row for row in snapshot["decision_history"]
              if row.get("checkpoint_id") == checkpoint["checkpoint_id"]]
    assert [row["event_type"] for row in events] == [
        "use_checkpoint_created", "use_checkpoint_metadata_corrected",
        "use_checkpoint_withdrawn"]
