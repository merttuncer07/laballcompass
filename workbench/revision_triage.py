"""Deterministic reviewer triage for confirmed workbook revisions."""
from __future__ import annotations

from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter


CHANGE_CATEGORIES = (
    "value_change",
    "formula_logic_change",
    "formula_to_literal",
    "literal_to_formula",
    "added_content",
    "removed_content",
    "movement_candidate",
    "unresolved",
)


def _category(cell):
    comparison = cell.get("comparison")
    left_formula = cell.get("left_type") == "formula"
    right_formula = cell.get("right_type") == "formula"
    if comparison == "changed_formula":
        if left_formula and not right_formula:
            return "formula_to_literal", None
        if right_formula and not left_formula:
            return "literal_to_formula", None
        return "formula_logic_change", None
    if comparison == "changed_value":
        return "value_change", None
    if comparison == "added":
        return "added_content", "formula_introduced" if right_formula else None
    if comparison == "removed":
        return "removed_content", "formula_removed" if left_formula else None
    if comparison == "uncompared":
        return "unresolved", None
    return None, None


def _point(cell, category):
    coordinate = cell["right"] if category == "added_content" else cell["left"]
    return coordinate_to_tuple(coordinate)


def _components(cells):
    by_point = {_point(cell, cell["_category"]): cell for cell in cells}
    pending = set(by_point)
    while pending:
        start = min(pending)
        pending.remove(start)
        component, todo = [by_point[start]], [start]
        while todo:
            row, column = todo.pop()
            # Reuse the content engine's established one-cell-gap neighborhood:
            # nearby changes may form one review region without claiming meaning.
            for dr in range(-2, 3):
                for dc in range(-2, 3):
                    neighbor = row + dr, column + dc
                    if neighbor in pending:
                        pending.remove(neighbor)
                        todo.append(neighbor)
                        component.append(by_point[neighbor])
        yield component


def _bounds(cells, side):
    coordinates = [cell[side] for cell in cells if cell.get(side)]
    if not coordinates:
        return None
    points = [coordinate_to_tuple(coordinate) for coordinate in coordinates]
    lo_row, hi_row = min(row for row, _ in points), max(row for row, _ in points)
    lo_col, hi_col = min(column for _, column in points), max(column for _, column in points)
    return f"{get_column_letter(lo_col)}{lo_row}:{get_column_letter(hi_col)}{hi_row}"


def _group_id(group):
    identity = json.dumps({key: group[key] for key in (
        "category", "left_sheet", "right_sheet", "left_range", "right_range", "changed_cell_count",
    )}, sort_keys=True)
    return "triage_" + hashlib.sha256(identity.encode()).hexdigest()[:16]


def _observations(cells, limit=5):
    ordered = sorted(cells, key=lambda cell: (_point(cell, cell["_category"]), cell["left"], cell["right"]))
    return [{key: cell.get(key) for key in (
        "left", "right", "left_value", "right_value", "left_type", "right_type", "comparison"
    )} for cell in ordered[:limit]]


def _make_group(sheet, category, cells, transition=None, movement=False, evidence=None,
                right_sheet=None):
    targets = sorted({target for cell in cells for target in cell.get("resolved_downstream_targets", [])})
    unresolved_cells = sorted({node for cell in cells for node in cell.get("unresolved_dependency_cells", [])})
    reasons = []
    if category == "unresolved":
        reasons.append("One or more compared positions could not be interpreted safely.")
    if unresolved_cells:
        reasons.append("A changed formula cell is outside the fully resolved dependency graph.")
    group = {
        "category": category,
        "sheet": sheet,
        "left_sheet": sheet,
        "right_sheet": right_sheet or sheet,
        "left_range": _bounds(cells, "left"),
        "right_range": _bounds(cells, "right"),
        "changed_cell_count": len(cells),
        "representative_observations": _observations(cells),
        "formula_changed": category in (
            "formula_logic_change", "formula_to_literal", "literal_to_formula"
        ) or transition in ("formula_introduced", "formula_removed"),
        "formula_transition": transition,
        "structural_movement_detected": movement,
        "movement_evidence": evidence,
        "resolved_downstream_target_count": len(targets),
        "resolved_downstream_targets": targets,
        "unresolved_coverage": bool(reasons),
        "unresolved_reasons": reasons,
        "unresolved_dependency_cells": unresolved_cells,
        "raw_cells": [{"left": cell.get("left"), "right": cell.get("right"),
                       "comparison": cell.get("comparison")} for cell in cells],
    }
    group["id"] = _group_id(group)
    return group


def _movement_groups(structural, raw_by_sheet):
    groups, masked = [], defaultdict(set)
    if not structural:
        return groups, masked
    for block in structural.get("blocks", []):
        if block.get("kind") != "row_alignment":
            continue
        moved = []
        differences = []
        impact_by_pair = {(row["left"], row["right"]): row
                          for row in block.get("cell_difference_impacts", [])}
        for original in block.get("cells", []):
            if not original.get("equal"):
                cell = dict(original)
                left_formula = cell.get("left_type") == "formula"
                right_formula = cell.get("right_type") == "formula"
                if left_formula and not right_formula:
                    category = "formula_to_literal"
                elif right_formula and not left_formula:
                    category = "literal_to_formula"
                elif left_formula and right_formula:
                    category = "formula_logic_change"
                elif "uncompared" in (cell.get("left_type"), cell.get("right_type")):
                    category = "unresolved"
                else:
                    category = "value_change"
                cell["comparison"] = "changed_formula" if "formula" in category else (
                    "uncompared" if category == "unresolved" else "changed_value")
                cell["_category"] = category
                impact = impact_by_pair.get((cell["left"], cell["right"]), {})
                cell["resolved_downstream_targets"] = impact.get("resolved_downstream_targets", [])
                cell["unresolved_dependency_cells"] = impact.get("unresolved_dependency_cells", [])
                differences.append(cell)
                continue
            if original.get("left") == original.get("right"):
                continue
            cell = dict(original)
            cell["comparison"] = "movement_candidate"
            cell["_category"] = "movement_candidate"
            moved.append(cell)
        if not moved:
            moved = []
        left_sheet, right_sheet = block["left"]["sheet"], block["right"]["sheet"]
        covered = set()
        for cell in moved:
            if cell["left"] in raw_by_sheet.get(left_sheet, set()):
                masked[left_sheet].add(cell["left"]); covered.add((left_sheet, cell["left"]))
            if cell["right"] in raw_by_sheet.get(right_sheet, set()):
                masked[right_sheet].add(cell["right"]); covered.add((right_sheet, cell["right"]))
        for cell in differences:
            if cell["left"] in raw_by_sheet.get(left_sheet, set()):
                masked[left_sheet].add(cell["left"])
            if cell["right"] in raw_by_sheet.get(right_sheet, set()):
                masked[right_sheet].add(cell["right"])
        alignment_evidence = {
            "kind": "row_alignment",
            "matching_rows": block.get("matching_rows", 0),
            "matching_cells": block.get("matching_cells", 0),
            "compared_positions": block.get("compared_positions", 0),
            "raw_changed_positions_normalized": len(covered),
            "column_mapping": block.get("column_mapping", []),
            "row_mapping": block.get("row_mapping", []),
            "interpretation": block.get("interpretation"),
        }
        if covered:
            group = _make_group(
                left_sheet, "movement_candidate", moved, movement=True,
                right_sheet=right_sheet,
                evidence=alignment_evidence)
            targets = sorted(set(block.get("formula_targets", [])))
            group["resolved_downstream_targets"] = targets
            group["resolved_downstream_target_count"] = len(targets)
            group["changed_cell_count"] = len(covered)
            group["id"] = _group_id(group)
            groups.append(group)
        by_category = defaultdict(list)
        for cell in differences:
            by_category[cell["_category"]].append(cell)
        for category, cells in sorted(by_category.items()):
            groups.append(_make_group(left_sheet, category, cells, movement=True,
                                      evidence={**alignment_evidence,
                                                "kind": "row_alignment_difference"},
                                      right_sheet=right_sheet))
    return groups, masked


def build_revision_triage(exact, structural=None):
    """Collapse raw same-layout changes into deterministic reviewer groups."""
    if exact.get("method") != "same_layout_v1":
        raise ValueError("Revision triage requires an exact same-layout comparison")
    raw = []
    raw_by_sheet = defaultdict(set)
    for block in exact.get("blocks", []):
        sheet = block["left"]["sheet"]
        impacts = {(row["left"], row["right"]): row
                   for row in block.get("cell_difference_impacts", [])}
        for original in block.get("cells", []):
            category, transition = _category(original)
            if category is None:
                continue
            cell = dict(original)
            cell["_sheet"] = sheet
            cell["_category"] = category
            cell["_transition"] = transition
            impact = impacts.get((cell["left"], cell["right"]), {})
            cell["resolved_downstream_targets"] = impact.get("resolved_downstream_targets", [])
            cell["unresolved_dependency_cells"] = impact.get("unresolved_dependency_cells", [])
            raw.append(cell)
            raw_by_sheet[sheet].add(cell["left"])
            raw_by_sheet[sheet].add(cell["right"])

    movement, masked = _movement_groups(structural, raw_by_sheet)
    grouped = defaultdict(list)
    for cell in raw:
        if cell["left"] in masked.get(cell["_sheet"], set()) or cell["right"] in masked.get(cell["_sheet"], set()):
            continue
        grouped[cell["_sheet"], cell["_category"], cell["_transition"]].append(cell)

    groups = list(movement)
    for (sheet, category, transition), cells in sorted(
            grouped.items(), key=lambda item: (item[0][0], item[0][1], item[0][2] or "")):
        for component in _components(cells):
            groups.append(_make_group(sheet, category, component, transition=transition))

    for sheet in exact.get("unmatched_sheets", {}).get("left", []):
        placeholder = {"left": "A1", "right": None, "left_value": "Entire sheet present only before",
                       "right_value": None, "left_type": "sheet", "right_type": "missing",
                       "comparison": "removed", "_category": "removed_content"}
        group = _make_group(sheet, "removed_content", [placeholder])
        group["changed_cell_count"] = 0
        group["raw_cells"] = []
        group["unresolved_coverage"] = True
        group["unresolved_reasons"] = ["The unmatched sheet was not compared cell by cell."]
        group["id"] = _group_id(group)
        groups.append(group)
    for sheet in exact.get("unmatched_sheets", {}).get("right", []):
        placeholder = {"left": None, "right": "A1", "left_value": None,
                       "right_value": "Entire sheet present only after", "left_type": "missing",
                       "right_type": "sheet", "comparison": "added", "_category": "added_content"}
        group = _make_group(sheet, "added_content", [placeholder])
        group["changed_cell_count"] = 0
        group["raw_cells"] = []
        group["unresolved_coverage"] = True
        group["unresolved_reasons"] = ["The unmatched sheet was not compared cell by cell."]
        group["id"] = _group_id(group)
        groups.append(group)

    groups.sort(key=lambda group: (
        -group["resolved_downstream_target_count"], group["category"], group["sheet"],
        group["left_range"] or "", group["right_range"] or "", group["id"],
    ))
    distribution = dict(sorted(Counter(group["category"] for group in groups).items()))
    raw_count = len(raw)
    group_count = len(groups)
    coverage = exact.get("dependency_coverage", {
        "lineage_available": exact.get("lineage_available", False),
        "unresolved_formula_count": None,
    })
    for group in groups:
        group["dependency_coverage"] = {
            "lineage_available": coverage.get("lineage_available", False),
            "direct_unresolved": group["unresolved_coverage"],
            "workbook_unresolved_formula_count": coverage.get("unresolved_formula_count"),
        }
        if not coverage.get("lineage_available", False) and not group["unresolved_coverage"]:
            group["unresolved_coverage"] = True
            group["unresolved_reasons"].append("Structural reach is unavailable for this workbook pair.")
    metrics = {
        "raw_change_count": raw_count,
        "triage_group_count": group_count,
        "compression_ratio": round(raw_count / group_count, 6) if group_count else None,
        "groups_with_resolved_impact": sum(group["resolved_downstream_target_count"] > 0 for group in groups),
        "groups_without_resolved_impact": sum(group["resolved_downstream_target_count"] == 0 for group in groups),
        "groups_with_unresolved_coverage": sum(group["unresolved_coverage"] for group in groups),
        "formula_logic_change_groups": distribution.get("formula_logic_change", 0),
        "formula_to_literal_groups": distribution.get("formula_to_literal", 0),
        "literal_to_formula_groups": distribution.get("literal_to_formula", 0),
        "movement_candidate_groups": distribution.get("movement_candidate", 0),
    }
    reaches = [group["resolved_downstream_target_count"] for group in groups]
    reach_distribution = {str(reach): count for reach, count in sorted(Counter(reaches).items())}
    return {
        "schema_version": 1,
        "method": "deterministic_revision_triage_v1",
        "input_files": exact.get("input_files", []),
        "summary": metrics,
        "group_type_distribution": distribution,
        "structural_reach": {
            "minimum": min(reaches) if reaches else 0,
            "maximum": max(reaches) if reaches else 0,
            "resolved_target_count_distribution": reach_distribution,
        },
        "dependency_coverage": coverage,
        "movement_analysis": {
            "available": structural is not None,
            "ambiguous_rows_omitted": (structural or {}).get("record_alignment", {}).get("ambiguous_rows_omitted", 0),
            "scope": (structural or {}).get("record_alignment", {}).get("scope"),
        },
        "groups": groups,
        "scope": [
            "Groups are deterministic geometric or correspondence-based review units; they do not assign audit meaning, risk or materiality.",
            "Structural reach counts only terminal formulas reachable in the supported static graph. Zero resolved reach does not prove no effect when dependency coverage is unresolved.",
            "Movement candidates use existing mutually unique record correspondence. Raw same-layout evidence remains available in the linked comparison.",
            "Formula text is compared without evaluating workbook results. Formatting and VBA are outside this analysis.",
        ],
    }


def write_revision_triage(exact_path, destination, structural_path=None):
    exact = json.loads(Path(exact_path).read_text())
    structural = json.loads(Path(structural_path).read_text()) if structural_path and Path(structural_path).is_file() else None
    result = build_revision_triage(exact, structural)
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    (destination / "triage.json").write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")
    from .revision_triage_report import write_revision_triage_report
    write_revision_triage_report(result, destination / "index.html")
    return destination / "index.html"
