"""Effective static ranges for Excel's single-condition aggregation functions."""
import re
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

FUNCTIONS = {'SUMIF', 'AVERAGEIF'}
CONTAINERS = {'FUNC', 'PAREN', 'ARRAY'}


def _calls(tokens):
    stack = []
    for index, token in enumerate(tokens):
        if token.type in CONTAINERS and token.subtype == 'OPEN':
            stack.append({'kind': token.type, 'name': token.value[:-1].upper().replace('_XLFN.', '').replace('_XLWS.', ''),
                          'start': index + 1, 'arguments': []})
        elif token.type == 'SEP' and token.subtype == 'ARG' and stack and stack[-1]['kind'] == 'FUNC':
            frame = stack[-1]
            frame['arguments'].append(range(frame['start'], index))
            frame['start'] = index + 1
        elif token.type in CONTAINERS and token.subtype == 'CLOSE':
            if not stack or stack[-1]['kind'] != token.type:
                raise ValueError('Unbalanced formula containers')
            frame = stack.pop()
            if frame['kind'] == 'FUNC' and frame['name'] in FUNCTIONS:
                frame['arguments'].append(range(frame['start'], index))
                yield frame['name'], frame['arguments']
    if stack:
        raise ValueError('Unclosed formula container')


def _single_reference(tokens, indices):
    indices = [i for i in indices if tokens[i].type != 'WHITE-SPACE']
    # Parentheses around a reference do not alter its rectangular extent.
    while len(indices) >= 3 and tokens[indices[0]].type == tokens[indices[-1]].type == 'PAREN' and tokens[indices[0]].subtype == 'OPEN' and tokens[indices[-1]].subtype == 'CLOSE':
        indices = indices[1:-1]
    if len(indices) != 1:
        raise ValueError('Range argument must resolve to one static rectangular reference')
    index = indices[0]
    if tokens[index].type != 'OPERAND' or tokens[index].subtype != 'RANGE':
        raise ValueError('Range argument must resolve to one static rectangular reference')
    return index, tokens[index].value


def _rectangle(nodes):
    if not nodes:
        raise ValueError('Empty range argument')
    origins = set()
    coordinates = set()
    for node in nodes:
        origin, address = node.rsplit('!', 1)
        if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', address):
            raise ValueError('Unsupported resolved coordinate')
        column, row = coordinate_from_string(address)
        column = column_index_from_string(column)
        if column > 16384 or row > 1048576:
            raise ValueError('Reference is outside Excel bounds')
        origins.add(origin)
        coordinates.add((column, row))
    if len(origins) != 1:
        raise ValueError('Range argument spans multiple worksheets or workbooks')
    left = min(c for c, _ in coordinates); right = max(c for c, _ in coordinates)
    top = min(r for _, r in coordinates); bottom = max(r for _, r in coordinates)
    if (right - left + 1) * (bottom - top + 1) != len(coordinates):
        raise ValueError('Range argument is not a single rectangle')
    return origins.pop(), left, top, right - left + 1, bottom - top + 1


def conditional_range_overrides(tokens, resolve, cell):
    """Return token-index overrides, observable range adjustments and unresolved reasons.

    resolve must reject multiple declared areas before flattening them to cells.
    No values, criteria matches, or workbook formulas are evaluated here.
    """
    overrides, adjustments, issues = {}, [], []
    for function, arguments in _calls(tokens):
        if len(arguments) == 2:
            continue
        if len(arguments) != 3:
            issues.append(function + ': expected two or three arguments')
            continue
        try:
            _, criteria_ref = _single_reference(tokens, arguments[0])
            value_index, value_ref = _single_reference(tokens, arguments[2])
            criteria = _rectangle(resolve(criteria_ref))
            declared_nodes = set(resolve(value_ref))
            origin, left, top, _, _ = _rectangle(declared_nodes)
            width, height = criteria[3:]
            if width * height > 10000:
                raise ValueError('Reference expansion exceeds 10,000 cells')
            right, bottom = left + width - 1, top + height - 1
            if right > 16384 or bottom > 1048576:
                raise ValueError('Effective range is outside Excel bounds')
            effective = {f'{origin}!{get_column_letter(c)}{r}' for c in range(left, right + 1) for r in range(top, bottom + 1)}
            overrides[value_index] = effective
            if effective != declared_nodes:
                adjustments.append({'cell': cell, 'function': function, 'criteria_reference': criteria_ref,
                                    'declared_reference': value_ref,
                                    'effective_reference': f'{origin}!{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}',
                                    'declared_cell_count': len(declared_nodes), 'effective_cell_count': len(effective)})
        except ValueError as error:
            issues.append(function + ': ' + str(error))
    return overrides, adjustments, issues
