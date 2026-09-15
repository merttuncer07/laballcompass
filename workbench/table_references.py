"""Resolve saved Excel table metadata into static cell references.

No table-size inference from nearby values, formula evaluation or refresh.
Invalid/missing selectors stay unresolved rather than becoming empty evidence.
"""
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple, get_column_letter


def split_sheet_reference(token):
    """Find ! outside table brackets/quoted sheet names (headers may contain !)."""
    depth=0;quoted=False;last=None;i=0
    while i<len(token):
        c=token[i]
        if c=="'":
            if depth:
                if i+1<len(token) and token[i+1] in "[]#'@":i+=2;continue
            elif quoted and i+1<len(token) and token[i+1]=="'":i+=2;continue
            else:quoted=not quoted
        elif not quoted:
            if c=='[':depth+=1
            elif c==']':depth-=1
            elif c=='!' and depth==0:last=i
        i+=1
    return (token[:last],token[last+1:]) if last is not None else (None,token)


def _parse_selector(text):
    """Return leaf specifiers/operators, preserving escapes until classification."""
    def group(start,depth=0):
        if depth>3 or start>=len(text) or text[start]!='[':raise ValueError('Invalid table selector')
        i=start+1
        while i<len(text) and text[i].isspace():i+=1
        prefix=[]
        if text[i:i+2]=='@[':prefix=[('leaf','#This Row')];i+=1
        if i<len(text) and text[i]=='[':
            tokens=prefix
            while True:
                child,i=group(i,depth+1);tokens.extend(child)
                while i<len(text) and text[i].isspace():i+=1
                if i>=len(text):raise ValueError('Unclosed table selector')
                if text[i]==']':return tokens,i+1
                if text[i] not in ',:':raise ValueError('Invalid table selector separator')
                tokens.append(('operator',text[i]));i+=1
                while i<len(text) and text[i].isspace():i+=1
        # A leaf header may contain commas/colons; only unescaped brackets split it.
        i=start+1;value=''
        while i<len(text):
            c=text[i]
            if c=="'" and i+1<len(text) and text[i+1] in "[]#'@":
                value+=text[i:i+2];i+=2;continue
            if c==']':
                if not value:raise ValueError('Empty nested table selector')
                return [('leaf',value)],i+1
            if c=='[':raise ValueError('Unescaped bracket in table column')
            value+=c;i+=1
        raise ValueError('Unclosed table selector')
    if text in ('','[]'):return []
    tokens,end=group(0)
    if text[end:].strip():raise ValueError('Unsupported suffix after table selector')
    return tokens


def _unescape(value):
    out='';i=0
    while i<len(value):
        if value[i]=="'" and i+1<len(value) and value[i+1] in "[]#'@":i+=1
        out+=value[i];i+=1
    return out


class TableReferences:
    def __init__(self,books):
        self.tables=[]
        for workbook in books:
            lookup={}
            for sheet in workbook:
                for table in sheet.tables.values():
                    name=table.displayName.casefold()
                    lookup.setdefault(name,[]).append((sheet.title,table))
            self.tables.append(lookup)

    def resolve(self,address,book_id,current_sheet,current_coordinate):
        """Return (table sheet, coordinates) or None for an ordinary named range."""
        name,separator,rest=address.partition('[')
        if not separator and name.casefold() not in self.tables[book_id]:return None
        selector='['+rest if separator else ''
        if name:
            choices=self.tables[book_id].get(name.casefold(),[])
            if len(choices)!=1:raise ValueError('Unknown or ambiguous table: '+name)
            sheet,table=choices[0]
        else:
            if current_coordinate is None:raise ValueError('Unqualified table reference requires a formula location')
            row,col=coordinate_to_tuple(current_coordinate)
            choices=[]
            for entries in self.tables[book_id].values():
                for sheet,table in entries:
                    c1,r1,c2,r2=range_boundaries(table.ref)
                    if sheet==current_sheet and r1<=row<=r2 and c1<=col<=c2:choices.append((sheet,table))
            if len(choices)!=1:raise ValueError('Unqualified table reference is outside one unambiguous table')
            sheet,table=choices[0]
        c1,r1,c2,r2=range_boundaries(table.ref)
        headers=1 if table.headerRowCount is None else table.headerRowCount
        totals=table.totalsRowCount or 0
        if headers not in (0,1) or totals not in (0,1) or r2-r1+1<headers+totals:
            raise ValueError('Unsupported table row metadata')
        columns={}
        for i,column in enumerate(table.tableColumns):columns.setdefault(column.name.casefold(),[]).append(c1+i)
        if len(table.tableColumns)!=c2-c1+1 or any(len(v)!=1 for v in columns.values()):
            raise ValueError('Table column metadata is missing or ambiguous')
        tokens=_parse_selector(selector);selectors=[];colnames=[];operators=[]
        for kind,token in tokens:
            if kind=='operator':operators.append(token);continue
            if not token:continue
            if token.startswith('@'):
                selectors.append('#this row');token=token[1:]
                if not token:continue
            if token.startswith('#'):selectors.append(token.casefold())
            else:colnames.append(_unescape(token))
        if not set(selectors)<={'#all','#data','#headers','#totals','#this row'}:
            raise ValueError('Unknown table item selector')
        if len(selectors)!=len(set(selectors)):raise ValueError('Repeated table item selector')
        if '#this row' in selectors and len(selectors)>1 or '#all' in selectors and len(selectors)>1:
            raise ValueError('Incompatible table item selectors')
        if len(selectors)>1 and set(selectors) not in ({'#headers','#data'},{'#data','#totals'}):
            raise ValueError('Unsupported combination of table row selectors')
        if len(colnames)>2 or operators.count(':')>1 or (len(colnames)==2 and operators.count(':')!=1) or (':' in operators and len(colnames)!=2):
            raise ValueError('Unsupported table column union/range')
        if colnames:
            try:chosen=[columns[n.casefold()][0] for n in colnames]
            except KeyError as error:raise ValueError('Unknown table column: '+str(error)) from error
            left,right=min(chosen),max(chosen)
        else:left,right=c1,c2
        selectors=set(selectors) or {'#data'}
        data_start,data_end=r1+headers,r2-totals
        selected_rows=set()
        def add_rows(start,end):
            if (end-start+1)*(right-left+1)>10000:
                raise ValueError('Table reference expansion exceeds 10,000 cells')
            selected_rows.update(range(start,end+1))
        for item in selectors:
            if item=='#all':add_rows(r1,r2)
            elif item=='#headers':
                if not headers:raise ValueError('Table has no header row')
                selected_rows.add(r1)
            elif item=='#totals':
                if not totals:raise ValueError('Table has no totals row')
                selected_rows.add(r2)
            elif item=='#this row':
                if current_coordinate is None:raise ValueError('This Row requires a formula location')
                row,_=coordinate_to_tuple(current_coordinate)
                if not data_start<=row<=data_end:raise ValueError('This Row is outside table data rows')
                selected_rows.add(row)
            else:add_rows(data_start,data_end)
        if not selected_rows:raise ValueError('Table reference has no data rows')
        if len(selected_rows)*(right-left+1)>10000:raise ValueError('Table reference expansion exceeds 10,000 cells')
        return sheet,{f'{get_column_letter(col)}{row}' for row in selected_rows for col in range(left,right+1)}
