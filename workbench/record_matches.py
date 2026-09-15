"""Locate reordered records using observed column values and row agreement.

Column bags propose an alignment; record-level agreement verifies it. Equal
marginal distributions alone do not create a match. Ambiguous column/row ties
are omitted instead of inventing a specific correspondence.
"""
from collections import defaultdict, Counter
from itertools import combinations


def _useful(token):
    return token[0]!='bool' and token not in (('number',0),('number',1),('number',-1))


def find_record_matches(sheets):
    from openpyxl.utils.cell import get_column_letter
    columns=[];index=defaultdict(list)
    for sid,sheet in enumerate(sheets):
        values=defaultdict(dict)
        for (row,col),token in sheet['values'].items():values[col][row]=token
        for col,cells in sorted(values.items()):
            tokens={t for t in cells.values() if _useful(t)}
            if len(tokens)<3:continue
            cid=len(columns);columns.append((sid,col,cells,tokens))
            for token in tokens:index[token].append(cid)
    overlaps=Counter();comparisons=0
    for token,cids in index.items():
        if len(cids)>64:continue
        for a,b in combinations(cids,2):
            if columns[a][0]==columns[b][0]:continue
            overlaps[a,b]+=1;comparisons+=1
            if comparisons>1000000:
                raise ValueError('Record alignment exceeds 1,000,000 column candidate comparisons; supply a smaller collection')
    proposals=defaultdict(dict)
    for (a,b),count in overlaps.items():
        if count<3:continue
        ta,tb=columns[a][3],columns[b][3]
        if count/min(len(ta),len(tb))<.6:continue
        # Jaccard prefers a specific match over an arbitrary large superset.
        score=count/len(ta|tb)
        proposals[columns[a][0],columns[b][0]][a,b]=score
    result=[];row_comparisons=0;details=0;ambiguous=0;omitted=0
    coordinate=lambda r,c:get_column_letter(c)+str(r)
    for (sa,sb),scores in sorted(proposals.items()):
        best_a=defaultdict(list);best_b=defaultdict(list)
        for (a,b),score in scores.items():best_a[a].append((score,b));best_b[b].append((score,a))
        def unique_best(items):
            top=max(s for s,_ in items)
            wins=[other for score,other in items if abs(score-top)<1e-12]
            return wins[0] if len(wins)==1 else None
        mapping=[]
        for a,choices in best_a.items():
            b=unique_best(choices)
            if b is not None and unique_best(best_b[b])==a:mapping.append((a,b))
        mapping.sort(key=lambda pair:columns[pair[0]][1])
        if len(mapping)<2:continue
        row_votes=defaultdict(set)
        for a,b in mapping:
            left,right=columns[a][2],columns[b][2]
            inverted_a=defaultdict(list);inverted_b=defaultdict(list)
            for row,token in left.items():
                if _useful(token):inverted_a[token].append(row)
            for row,token in right.items():
                if _useful(token):inverted_b[token].append(row)
            for token in inverted_a.keys()&inverted_b.keys():
                ra,rb=inverted_a[token],inverted_b[token]
                if len(ra)*len(rb)>4096:continue
                for x in ra:
                    for y in rb:
                        row_votes[x,y].add((a,token));row_comparisons+=1
                        if row_comparisons>1000000:
                            raise ValueError('Record alignment exceeds 1,000,000 row candidate comparisons; supply a smaller collection')
        candidates={}
        for (ra,rb),votes in row_votes.items():
            if len({t for a,t in votes})<2:continue
            equal=sum(columns[a][2].get(ra) is not None and columns[a][2].get(ra)==columns[b][2].get(rb) for a,b in mapping)
            if equal/len(mapping)<.8:continue
            candidates[ra,rb]=equal
        left_best=defaultdict(list);right_best=defaultdict(list)
        for (ra,rb),score in candidates.items():left_best[ra].append((score,rb));right_best[rb].append((score,ra))
        row_mapping=[]
        for ra,choices in sorted(left_best.items()):
            rb=unique_best(choices)
            if rb is not None and unique_best(right_best[rb])==ra:row_mapping.append((ra,rb))
            else:ambiguous+=1
        if len(row_mapping)<3:continue
        if details+len(row_mapping)*len(mapping)>100000:
            omitted+=1
            continue
        left,right=sheets[sa],sheets[sb];cells=[]
        for ra,rb in row_mapping:
            for a,b in mapping:
                ca,cb=columns[a][1],columns[b][1];pa,pb=(ra,ca),(rb,cb)
                va,vb=left['values'].get(pa),right['values'].get(pb)
                cells.append({'left':coordinate(*pa),'right':coordinate(*pb),
                              'left_value':left['raw'].get(pa),'right_value':right['raw'].get(pb),
                              'left_type':va[0] if va else 'uncompared','right_type':vb[0] if vb else 'uncompared',
                              'equal':va is not None and va==vb})
        matching=sum(c['equal'] for c in cells)
        unmatched_rows=[];unmatched_columns=[]
        for side,sheet in enumerate((left,right)):
            used_rows={pair[side] for pair in row_mapping}
            used_cols={columns[pair[side]][1] for pair in mapping}
            unmatched_rows.append(sorted({r for r,c in sheet['values']}-used_rows))
            unmatched_columns.append([get_column_letter(c) for c in sorted({c for r,c in sheet['values']}-used_cols)])
        def location(sheet,side):
            rr=[pair[side] for pair in row_mapping];cc=[columns[pair[side]][1] for pair in mapping]
            return {'workbook':sheet['workbook'],'sheet':sheet['sheet'],
                    'range':coordinate(min(rr),min(cc))+':'+coordinate(max(rr),max(cc))}
        result.append({'kind':'row_alignment','left':location(left,0),'right':location(right,1),
                       'matching_rows':len(row_mapping),'row_mapping':[{'left':a,'right':b} for a,b in row_mapping],
                       'unmapped_left_rows':unmatched_rows[0],'unmapped_right_rows':unmatched_rows[1],
                       'unmapped_left_columns':unmatched_columns[0],'unmapped_right_columns':unmatched_columns[1],
                       'column_mapping':[{'left':get_column_letter(columns[a][1]),'right':get_column_letter(columns[b][1])} for a,b in mapping],
                       'matching_cells':matching,'compared_positions':len(cells),'match_fraction':matching/len(cells),
                       'different_or_uncompared_cells':len(cells)-matching,'cells':cells,
                       'interpretation':'Observed row/value correspondences, not inferred copying or common origin. Only mapped rows and columns are shown; the bounding ranges may contain omitted cells.'})
        details+=len(cells)
    return result,{'column_candidate_comparisons':comparisons,'row_candidate_comparisons':row_comparisons,
                   'ambiguous_rows_omitted':ambiguous,
                   'regions_omitted_due_detail_limit':omitted,
                   'scope':'Mutual unique best column value-set alignment, then at least 3 mutually unambiguous rows with 80% mapped-column agreement and 2 distinct nontrivial anchor values. Unmapped columns and ambiguous rows are omitted.'}
