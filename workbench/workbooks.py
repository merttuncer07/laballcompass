"""Conservative static XLSX lineage. Never computes or certifies spreadsheet values."""
from collections import defaultdict, deque
from contextlib import ExitStack, contextmanager
import json
from io import BytesIO
from urllib.parse import unquote, urlsplit
import hashlib
from pathlib import Path
import re
import zipfile
from .contracts import DependencyProblem
from .function_references import conditional_range_overrides
from .table_references import TableReferences, split_sheet_reference
from .limits import (
    LINEAGE_NODES,
    POPULATED_CELLS_PER_COLLECTION,
    POPULATED_CELLS_PER_WORKBOOK,
    WORKSHEET_BOUNDING_CELLS,
)

DIRECT = re.compile(r'^\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?$')
DYNAMIC = {'INDIRECT', 'OFFSET', 'LAMBDA', 'LET'}
KNOWN_FUNCTIONS = {'SUM', 'AVERAGE', 'MIN', 'MAX', 'COUNT', 'COUNTA', 'COUNTBLANK', 'COUNTIF', 'COUNTIFS',
                   'SUMIF', 'SUMIFS', 'AVERAGEIF', 'AVERAGEIFS', 'IF', 'IFS', 'IFERROR', 'IFNA',
                   'AND', 'OR', 'NOT', 'XOR', 'ABS', 'ROUND', 'ROUNDUP', 'ROUNDDOWN', 'INT', 'MOD',
                   'POWER', 'SQRT', 'PRODUCT', 'SUMPRODUCT', 'SUBTOTAL', 'AGGREGATE',
                   'INDEX', 'MATCH', 'VLOOKUP', 'HLOOKUP', 'XLOOKUP', 'XMATCH', 'ROWS', 'COLUMNS',
                   'CONCAT', 'CONCATENATE', 'TEXTJOIN', 'TEXT', 'LEFT', 'RIGHT', 'MID', 'LEN',
                   'TRIM', 'UPPER', 'LOWER', 'VALUE', 'ISNUMBER', 'ISTEXT', 'ISBLANK', 'ISERROR',
                   'ISNA', 'DATE', 'YEAR', 'MONTH', 'DAY', 'DAYS', 'TRUE', 'FALSE'}


def formula_descriptor(cell):
    """Return stable text for ordinary and Excel array formulas."""
    if cell.data_type != 'f':
        return None
    value = cell.value
    if isinstance(value, str):
        return value
    text = getattr(value, 'text', None)
    if not isinstance(text, str):
        return None
    return text


def analyze_workbook(path):
    return analyze_workbooks([path])


def analyze_workbooks(paths):
    with open_workbooks(paths) as (paths, books):
        return _analyze_loaded(paths, books)


class LineageUnavailable(ValueError):
    """Valid workbook input without a resolved formula graph."""

    def __init__(self, message, *, external_dependencies=None):
        super().__init__(message)
        self.external_dependencies = external_dependencies or {
            "parser_available": False, "references": [],
            "limitations": ["Formula lineage was unavailable before external references could be summarized."],
        }


def _external_dependency_observations(paths, linked_workbooks, issues):
    """Expose conservative observations already produced by the formula parser."""
    references = []
    for link in linked_workbooks.values():
        locator = str(link["declared_locator"])
        filename = unquote(urlsplit(locator.replace('\\', '/')).path).rstrip('/').rsplit('/', 1)[-1]
        references.append({
            "from_workbook": link["from_workbook"],
            "formula_cell": None,
            "reference": filename or locator,
            "declared_locator": locator,
            "status": "resolved_against_supplied_evidence",
            "matched_workbook": link["matched_workbook"],
            "identity_basis": "Unique supplied filename match",
        })
    prefixes = (
        "External workbook not supplied: ",
        "Unresolved external workbook index: ",
        "Unsupported external workbook reference: ",
    )
    multiple = len(paths) > 1
    for cell, reasons in sorted(issues.items()):
        from_workbook = cell.split("::", 1)[0] if multiple and "::" in cell else paths[0].name
        for reason in reasons:
            prefix = next((value for value in prefixes if reason.startswith(value)), None)
            if prefix is None:
                continue
            reference = reason[len(prefix):]
            known_identity = prefix == "External workbook not supplied: "
            references.append({
                "from_workbook": from_workbook,
                "formula_cell": cell,
                "reference": reference,
                "declared_locator": None,
                "status": "unresolved_or_missing",
                "matched_workbook": None,
                "identity_basis": ("Filename parsed from static formula reference" if known_identity
                                   else "Before/after identity cannot be established safely"),
            })
    unique = {}
    for row in references:
        key = (row["from_workbook"], row["formula_cell"], row["reference"], row["status"])
        unique[key] = row
    return {
        "parser_available": True,
        "references": [unique[key] for key in sorted(unique)],
        "limitations": [
            "Only static external references observed by the existing formula parser are included.",
            "A filename match does not establish workbook version or authenticity.",
        ],
    }


@contextmanager
def open_workbooks(paths, *, preserve_order=False):
    """Read only supplied files; external locators never trigger file/network access."""
    from openpyxl import load_workbook
    paths = list(dict.fromkeys(Path(p).resolve() for p in paths))
    if not preserve_order:
        paths.sort(key=lambda p: (p.name.casefold(), str(p)))
    if not 1 <= len(paths) <= 20:
        raise ValueError('Supply between 1 and 20 workbooks')
    if len({p.name.casefold() for p in paths}) != len(paths):
        raise ValueError('Ambiguous workbook names: supplied files must have distinct filenames')
    expanded = 0
    with ExitStack() as stack:
        books = []
        for path in paths:
            if path.suffix.lower() not in ('.xlsx', '.xlsm'):
                raise ValueError('The pilot reads .xlsx and .xlsm workbooks')
            with path.open('rb') as input_file:
                raw = input_file.read(32 * 1024**2 + 1)
            if len(raw) > 32 * 1024**2:
                raise ValueError('Pilot workbook limit: 32 MB compressed')
            stream = stack.enter_context(BytesIO(raw))
            with zipfile.ZipFile(stream) as archive:
                expanded += sum(i.file_size for i in archive.infolist())
            if expanded > 256 * 1024**2:
                raise ValueError('Pilot workbook collection limit: 256 MB expanded')
            stream.seek(0)
            workbook = load_workbook(stream, data_only=False, keep_links=True, keep_vba=False)
            workbook._lab_input_sha256 = hashlib.sha256(raw).hexdigest()
            stack.callback(workbook.close)
            books.append(workbook)
        yield paths, books


def _analyze_loaded(paths, books):
    from openpyxl.formula.tokenizer import Tokenizer
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    multiple = len(paths) > 1
    book_names = {p.name.casefold(): i for i, p in enumerate(paths)}
    linked_workbooks = {}
    tables = TableReferences(books)
    table_links = {}
    formulas = {}
    cells = {}
    populated_by_workbook = defaultdict(int)
    labels = {}
    locations = {}

    def key(book_id, sheet, coordinate):
        prefix = paths[book_id].name + '::' if multiple else ''
        return prefix + sheet + '!' + coordinate.replace('$', '').upper()

    for book_id, workbook in enumerate(books):
        for sheet in workbook:
            if sheet.max_row * sheet.max_column > WORKSHEET_BOUNDING_CELLS:
                raise ValueError(f'Pilot worksheet bounding range limit: {WORKSHEET_BOUNDING_CELLS:,} cells')
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None: continue
                    node = key(book_id, sheet.title, cell.coordinate)
                    cells[node] = cell.data_type
                    labels[node] = node
                    populated_by_workbook[paths[book_id].name] += 1
                    if populated_by_workbook[paths[book_id].name] > POPULATED_CELLS_PER_WORKBOOK:
                        raise ValueError(
                            f'Pilot workbook populated-cell limit: {POPULATED_CELLS_PER_WORKBOOK:,} ({paths[book_id].name})')
                    if len(cells) > POPULATED_CELLS_PER_COLLECTION:
                        raise ValueError(
                            f'Pilot collection populated-cell limit: {POPULATED_CELLS_PER_COLLECTION:,}')
                    if cell.data_type == 'f':
                        formulas[node] = formula_descriptor(cell)
                        locations[node] = (book_id, sheet.title, cell.coordinate)
    if not formulas:
        raise LineageUnavailable('No formulas found. Hardcoded equal values alone cannot establish shared provenance.')
    if len(cells) > POPULATED_CELLS_PER_COLLECTION:
        raise ValueError(f'Pilot populated-cell limit: {POPULATED_CELLS_PER_COLLECTION:,}')
    dependencies = {}
    issues = defaultdict(list)
    function_range_adjustments = []

    def reference(token, book_id, current_sheet, current_coordinate=None, names_seen=frozenset(), *, single_area=False):
        workbook = books[book_id]
        sheet_name, address = split_sheet_reference(token)
        if sheet_name is not None:
            if sheet_name.startswith("'") and sheet_name.endswith("'"):
                sheet_name = sheet_name[1:-1].replace("''", "'")
        else:
            sheet_name, address = current_sheet, token
        if '[' in sheet_name or ']' in sheet_name:
            match = re.fullmatch(r"(.*?)\[([^\]]+)\](.+)", sheet_name)
            if not match:
                raise ValueError('Unsupported external workbook reference: ' + token)
            location_prefix, book_reference, sheet_name = match.groups()
            locator = location_prefix + book_reference
            if book_reference.isdecimal():
                index = int(book_reference) - 1
                links = workbook._external_links
                if not 0 <= index < len(links) or links[index].file_link is None:
                    raise ValueError('Unresolved external workbook index: ' + book_reference)
                locator = links[index].file_link.Target
            filename = unquote(urlsplit(str(locator).replace('\\', '/')).path).rstrip('/').rsplit('/', 1)[-1]
            destination = book_names.get(filename.casefold())
            if destination is None:
                raise ValueError('External workbook not supplied: ' + filename)
            source_name, destination_name = paths[book_id].name, paths[destination].name
            linked_workbooks[(source_name, str(locator), destination_name)] = {
                'from_workbook': source_name, 'declared_locator': str(locator),
                'matched_workbook': destination_name,
                'resolution': 'Unique supplied filename match; workbook version and authenticity are not established.'}
            book_id = destination
            workbook = books[book_id]

        if ':' in sheet_name:
            if single_area:
                raise ValueError('3D references are not a single function range: ' + token)
            endpoints = sheet_name.split(':')
            by_name = {s.casefold(): i for i, s in enumerate(workbook.sheetnames)}
            if len(endpoints) != 2 or any(s.casefold() not in by_name for s in endpoints):
                raise ValueError('Unknown 3D worksheet endpoint: ' + token)
            first, last = (by_name[s.casefold()] for s in endpoints)
            if first > last or not DIRECT.fullmatch(address):
                raise ValueError('Unsupported 3D worksheet range: ' + token)
            results = set()
            for sheet in workbook.sheetnames[first:last + 1]:
                results.update(reference("'" + sheet.replace("'", "''") + "'!" + address, book_id, current_sheet, current_coordinate))
                if len(results) > 10000:
                    raise ValueError('Reference expansion exceeds 10,000 cells: ' + token)
            return results
        sheet_name = next((s for s in workbook.sheetnames if s.casefold() == sheet_name.casefold()), sheet_name)
        if sheet_name not in workbook.sheetnames:
            raise ValueError('Unknown worksheet: ' + sheet_name)
        if DIRECT.fullmatch(address):
            low_col, low_row, high_col, high_row = range_boundaries(address)
            if high_col < low_col or high_row < low_row:
                raise ValueError('Reversed cell range: ' + token)
            if high_col > 16384 or high_row > 1048576:
                raise ValueError('Reference is outside Excel bounds')
            if (high_col - low_col + 1) * (high_row - low_row + 1) > 10000:
                raise ValueError('Reference expansion exceeds 10,000 cells: ' + token)
            return {key(book_id, sheet_name, f'{get_column_letter(col)}{row}')
                    for col in range(low_col, high_col + 1) for row in range(low_row, high_row + 1)}
        table_reference = tables.resolve(address, book_id, sheet_name, current_coordinate)
        if table_reference is not None:
            table_sheet, coordinates = table_reference
            table_links[(book_id,token,current_coordinate)] = {'workbook':paths[book_id].name,'token':token,'formula_coordinate':current_coordinate,'sheet':table_sheet,'referenced_cells':len(coordinates)}
            return {key(book_id,table_sheet,c) for c in coordinates}
        if '[' in address or '#' in address or '@' in address:
            raise ValueError('Unsupported table, spill or implicit-intersection reference: ' + token)
        if address.casefold() in names_seen:
            raise ValueError('Recursive defined name: ' + address)
        def find_name(mapping):
            return next((v for k, v in mapping.items() if k.casefold() == address.casefold()), None)
        name = find_name(workbook[sheet_name].defined_names) or find_name(workbook.defined_names)
        if name is None: raise ValueError('Unresolved name or reference: ' + token)
        # Names are accepted only when they resolve to static coordinates.
        if name.type != 'RANGE': raise ValueError('Non-range defined name: ' + address)
        results = set()
        destinations = list(name.destinations)
        if single_area and len(destinations) != 1:
            raise ValueError('Multi-area defined name is not a single function range: ' + address)
        for destination_sheet, destination in destinations:
            results.update(reference("'" + destination_sheet.replace("'", "''") + "'!" + destination,
                                     book_id, current_sheet, current_coordinate, names_seen | {address.casefold()}, single_area=single_area))
        if not results: raise ValueError('Empty defined name: ' + address)
        return results

    for node, formula in formulas.items():
        refs = set()
        if formula is None:
            issues[node].append('Array/shared formula representation unsupported')
        else:
            try:
                tokens = Tokenizer(formula).items
                overrides, adjustments, range_issues = conditional_range_overrides(
                    tokens, lambda text: reference(text, *locations[node], single_area=True), node)
                function_range_adjustments.extend(adjustments)
                if range_issues:
                    issues[node].extend(range_issues)
                for token_index, token in enumerate(tokens):
                    if token.type == 'FUNC' and token.subtype == 'OPEN':
                        function = token.value[:-1].upper().replace('_XLFN.', '').replace('_XLWS.', '')
                        if function in DYNAMIC or function not in KNOWN_FUNCTIONS:
                            issues[node].append('Dynamic or unsupported function: ' + function)
                    elif token.type == 'OPERAND' and token.subtype == 'RANGE':
                        try: refs.update(overrides[token_index] if token_index in overrides else reference(token.value, *locations[node]))
                        except ValueError as error: issues[node].append(str(error))
                    elif token.type == 'OPERAND' and token.subtype == 'ERROR':
                        issues[node].append('Formula contains error: ' + token.value)
            except Exception as error:
                issues[node].append('Formula parsing failed: ' + type(error).__name__)
        for ref in refs:
            if cells.get(ref) == 'e': issues[node].append('References an error-valued cell: ' + ref)
        dependencies[node] = refs

    # Mark every node downstream of a cycle without misclassifying the cycle as an input.
    dependents = defaultdict(list)
    indegree = {}
    for node, refs in dependencies.items():
        formula_refs = refs & formulas.keys()
        indegree[node] = len(formula_refs)
        for ref in formula_refs: dependents[ref].append(node)
    ready = deque(sorted(n for n, degree in indegree.items() if degree == 0))
    order = []
    while ready:
        node = ready.popleft()
        order.append(node)
        for child in sorted(dependents[node]):
            indegree[child] -= 1
            if indegree[child] == 0: ready.append(child)
    pending = set(formulas) - set(order)
    for node in pending: issues[node].append('Cycle or dependency downstream of a cycle')
    blocked = set(issues)
    for node in order:
        if dependencies[node] & blocked:
            blocked.add(node)
            issues[node].append('Depends on an unresolved formula')
    safe = set(formulas) - blocked
    safe_order = [node for node in order if node in safe]
    roots = {ref for node in safe for ref in dependencies[node] if ref not in formulas}
    if len(roots) + len(safe) > LINEAGE_NODES:
        raise ValueError(f'Expanded lineage exceeds {LINEAGE_NODES:,} nodes')
    external_dependencies = _external_dependency_observations(paths, linked_workbooks, issues)
    if not safe:
        raise LineageUnavailable(
            'No formula has a fully resolved static reference graph. '
            + '; '.join(f'{n}: {issues[n][0]}' for n in list(issues)[:5]),
            external_dependencies=external_dependencies)
    used = {ref for node in safe for ref in dependencies[node] if ref in safe}
    targets = tuple(sorted(safe - used))
    # A bit identifies a root once; descendant sets share compact immutable integers.
    # Unlike materializing a frozenset per formula, memory scales with graph size
    # and root bits rather than repeated Python set entries.
    root_order = sorted(roots)
    source_masks = {root: 1 << i for i, root in enumerate(root_order)}
    for node in safe_order:
        mask = 0
        for premise in dependencies[node]:
            mask |= source_masks[premise]
        source_masks[node] = mask
    groups = defaultdict(list)
    for node in safe:
        if source_masks[node]: groups[source_masks[node]].append(node)
    def root_cells(mask):
        cells = []
        while mask:
            bit = mask & -mask
            cells.append(root_order[bit.bit_length() - 1])
            mask ^= bit
        return cells
    same_origins = [{'root_cells': root_cells(mask), 'formula_cells': sorted(nodes)}
                    for mask, nodes in groups.items() if len(nodes) > 1]
    same_origins.sort(key=lambda g: (-len(g['formula_cells']), g['formula_cells']))
    input_files = [{'name': p.name, 'sha256': book._lab_input_sha256} for p, book in zip(paths, books)]
    digest = (hashlib.sha256(json.dumps(input_files, sort_keys=True).encode()).hexdigest()
              if multiple else input_files[0]['sha256'])
    metadata = {'title': f'{len(paths)} workbooks: ' + ', '.join(p.name for p in paths) if multiple else paths[0].name,
                'input_sha256': digest, 'input_files': input_files,
                'linked_workbooks': list(linked_workbooks.values()),
                'external_dependencies': external_dependencies,
                'structured_references': list(table_links.values()),
                'function_range_adjustments': function_range_adjustments,
                'origin': 'Workbook formulas parsed automatically; workbook values were not recalculated.',
                'labels': labels, 'formula_text': {n: formulas[n] for n in sorted(safe)},
                'formula_count': len(formulas), 'resolved_formula_count': len(safe),
                'unresolved': [{'cell': n, 'reasons': sorted(set(reasons))} for n, reasons in sorted(issues.items())],
                'same_origin_groups': same_origins,
                'assumptions': ['All static references are treated as potential dependencies, including unused IF branches and algebraic cancellations.',
                                'Distinct root cells are not evidence of statistically independent sources.',
                                'External links resolve only to uniquely named, explicitly supplied files. Saved table metadata resolves supported structured references. Missing files, unsupported dynamic references and cyclic paths remain unresolved; no external file was fetched.',
                                'Matching a referenced filename does not prove that the supplied file is the original version or is authentic.',
                                'Switching a source off reports structural sensitivity, not a false numeric value, audit failure or opinion.']}
    problem = DependencyProblem('workbook_' + metadata['input_sha256'][:12], {r: r for r in sorted(roots)},
                                tuple((n, tuple(sorted(dependencies[n]))) for n in safe_order), targets,
                                'static_formula_lineage', metadata)
    return problem
